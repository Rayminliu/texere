"""scripts/edit.py 的编辑链路。

用 python-docx 现场造 docx，不需要 Word，秒级可跑。

守的是两条：
  1. 跨 run 替换要正确，且保住原有 run 结构与格式；
  2. 「只改你指定的地方」——含图段落不动、歧义锚点拒绝、页眉页脚只动指定节。
"""

import json
import os
import subprocess
import sys

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

KIT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EDIT = os.path.join(KIT, "scripts", "edit.py")


def run_edit(path, *args, expect=0):
    """跑 edit.py；expect=None 表示不检查退出码。"""
    env = dict(os.environ, PYTHONIOENCODING="utf-8")
    r = subprocess.run([sys.executable, EDIT, str(path), *args], capture_output=True, env=env)
    out = (r.stdout + r.stderr).decode("utf-8", "replace")
    if expect is not None:
        assert r.returncode == expect, out
    return out


def make_docx(path, with_graphic=True, rows=2):
    """造一个贴近真实 Word 产物的文档：数字会被切成单独的 run。"""
    doc = Document()
    p = doc.add_paragraph()
    r0 = p.add_run("投标有效期：自开标之日起")
    r0.bold = True
    p.add_run(" 90 ")
    p.add_run("日历天")

    if with_graphic:
        pg = doc.add_paragraph()
        pg.add_run("图 1-1 总体架构图")
        pg._p.append(OxmlElement("w:drawing"))  # 只要有这个节点就算含图，不必真嵌图片

    t = doc.add_table(rows=rows, cols=2)
    t.cell(0, 0).text = "条款"
    t.cell(0, 1).text = "响应"
    if rows > 1:
        t.cell(1, 0).text = "工期"
        t.cell(1, 1).text = "180 天"
    doc.save(str(path))
    return str(path)


# ---------------------------------------------------------------- 替换


def test_replace_spanning_runs_keeps_structure(tmp_path):
    """跨 3 个 run 的替换：文本对、run 数不变、首个 run 的加粗保留。"""
    f = make_docx(tmp_path / "a.docx")
    run_edit(f, "--replace", "自开标之日起 90 日历天=自开标之日起 120 日历天")

    p = Document(f).paragraphs[0]
    assert p.text == "投标有效期：自开标之日起 120 日历天"
    assert len(p.runs) == 3, "run 结构必须保住，否则整段格式就没了"
    assert p.runs[0].bold is True, "新文本落在首个 run 上，其加粗不能丢"


def test_replace_within_single_run(tmp_path):
    """命中完全落在一个 run 内时，只改这个 run，其它 run 一个字符都不动。"""
    f = make_docx(tmp_path / "a.docx")
    run_edit(f, "--replace", "日历天=个工作日")
    p = Document(f).paragraphs[0]
    assert p.text == "投标有效期：自开标之日起 90 个工作日"
    assert [r.text for r in p.runs] == ["投标有效期：自开标之日起", " 90 ", "个工作日"]


def test_replace_leaves_graphic_paragraph_alone(tmp_path):
    """契约：含图段落一律跳过，绝不因为改文字把图搞坏。"""
    f = make_docx(tmp_path / "a.docx")
    out = run_edit(f, "--replace", "总体架构图=被改掉的图注")
    assert "跳过 1 个含图段落" in out
    assert "总体架构图" in Document(f).paragraphs[1].text


def test_replace_covers_tables_by_default(tmp_path):
    f = make_docx(tmp_path / "a.docx")
    run_edit(f, "--replace", "180 天=170 天")
    assert Document(f).tables[0].cell(1, 1).text == "170 天"


def test_replace_multiple_rules_and_missing(tmp_path):
    f = make_docx(tmp_path / "a.docx")
    out = run_edit(f, "--replace", "工期=质保期", "--replace", "不存在的文字=X")
    assert "replace: 1 处" in out


def test_replace_no_hit_warns(tmp_path):
    f = make_docx(tmp_path / "a.docx")
    out = run_edit(f, "--replace", "绝不存在=Y")
    assert "[warn]" in out and "一条都没命中" in out


def test_replace_bad_rule_exits(tmp_path):
    f = make_docx(tmp_path / "a.docx")
    run_edit(f, "--replace", "没有等号", expect=1)


# ---------------------------------------------------------------- 插入 / 删除


def test_insert_after_unique_anchor(tmp_path):
    f = make_docx(tmp_path / "a.docx", with_graphic=False)
    run_edit(f, "--after", "投标有效期", "--text", "补充：本条为新增。")
    texts = [p.text for p in Document(f).paragraphs]
    assert "补充：本条为新增。" in texts
    assert (
        texts.index("补充：本条为新增。") == texts.index("投标有效期：自开标之日起 90 日历天") + 1
    )


def test_insert_before_anchor(tmp_path):
    f = make_docx(tmp_path / "a.docx", with_graphic=False)
    run_edit(f, "--before", "投标有效期", "--text", "（征求意见稿）")
    texts = [p.text for p in Document(f).paragraphs]
    assert texts.index("（征求意见稿）") == texts.index("投标有效期：自开标之日起 90 日历天") - 1


def test_insert_multiline_keeps_order(tmp_path):
    f = make_docx(tmp_path / "a.docx", with_graphic=False)
    run_edit(f, "--after", "投标有效期", "--text", r"第一行\n第二行")
    texts = [p.text for p in Document(f).paragraphs]
    assert texts.index("第一行") + 1 == texts.index("第二行"), "多行插入顺序不能反"


def test_ambiguous_anchor_is_rejected(tmp_path):
    """目录刷成静态文本后，标题在目录与正文各一份——必须拒绝而不是两处都改。"""
    f = make_docx(tmp_path / "a.docx", with_graphic=False)
    doc = Document(f)
    doc.add_paragraph("1.2 资质与业绩")
    doc.add_paragraph("1.2 资质与业绩")
    doc.save(f)

    out = run_edit(f, "--after", "1.2 资质与业绩", "--text", "新增", expect=1)
    assert "命中 2 段" in out and "--all-anchors" in out
    assert "新增" not in [p.text for p in Document(f).paragraphs], "被拒绝时不能写盘"


def test_all_anchors_force(tmp_path):
    f = make_docx(tmp_path / "a.docx", with_graphic=False)
    doc = Document(f)
    doc.add_paragraph("重复标题")
    doc.add_paragraph("重复标题")
    doc.save(f)

    run_edit(f, "--after", "重复标题", "--text", "新增", "--all-anchors")
    assert [p.text for p in Document(f).paragraphs].count("新增") == 2


def test_delete_paragraph(tmp_path):
    f = make_docx(tmp_path / "a.docx", with_graphic=False)
    run_edit(f, "--delete", "投标有效期")
    assert "投标有效期" not in [p.text for p in Document(f).paragraphs]


def test_unknown_style_exits(tmp_path):
    f = make_docx(tmp_path / "a.docx", with_graphic=False)
    run_edit(f, "--after", "投标有效期", "--text", "X", "--style", "不存在的样式", expect=1)


# ---------------------------------------------------------------- 表格


def test_cell_set(tmp_path):
    f = make_docx(tmp_path / "a.docx")
    run_edit(f, "--cell", "0", "1", "1", "170 天")
    assert Document(f).tables[0].cell(1, 1).text == "170 天"


def test_cell_out_of_range_exits(tmp_path):
    f = make_docx(tmp_path / "a.docx")
    run_edit(f, "--cell", "0", "9", "1", "X", expect=1)


def test_add_and_del_row(tmp_path):
    f = make_docx(tmp_path / "a.docx")
    run_edit(f, "--add-row", "0", "质保期", "3 年")
    t = Document(f).tables[0]
    assert len(t.rows) == 3 and t.cell(2, 0).text == "质保期"

    run_edit(f, "--del-row", "0", "2")
    assert len(Document(f).tables[0].rows) == 2


# ---------------------------------------------------------------- 页眉页脚 / 分节


def test_header_only_touches_body_section(tmp_path):
    """封面节不该被加页眉。默认 --section body 只动最后一个节。"""
    f = tmp_path / "s.docx"
    doc = Document()
    doc.add_paragraph("第一页")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("第二页")
    doc.save(str(f))

    run_edit(f, "--header", "正文页眉")
    d = Document(f)
    assert d.sections[1].header.paragraphs[0].text == "正文页眉"
    assert d.sections[0].header.paragraphs[0].text != "正文页眉", "封面节被误加页眉"

    run_edit(f, "--footer", "— X —")
    assert Document(f).sections[1].footer.paragraphs[0].text == "— X —"


def test_header_all_sections(tmp_path):
    f = tmp_path / "s.docx"
    doc = Document()
    doc.add_paragraph("第一页")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("第二页")
    doc.save(str(f))

    run_edit(f, "--header", "全局页眉", "--section", "all")
    d = Document(f)
    assert d.sections[0].header.paragraphs[0].text == "全局页眉"
    assert d.sections[1].header.paragraphs[0].text == "全局页眉"


def test_bad_section_exits(tmp_path):
    f = make_docx(tmp_path / "a.docx")
    run_edit(f, "--header", "X", "--section", "99", expect=1)


# ---------------------------------------------------------------- 其他


def test_backup_is_written_by_default(tmp_path):
    f = make_docx(tmp_path / "a.docx")
    run_edit(f, "--replace", "工期=质保期")
    assert (tmp_path / "a.bak.docx").exists()
    assert Document(str(tmp_path / "a.bak.docx")).tables[0].cell(1, 0).text == "工期"


def test_no_backup(tmp_path):
    f = make_docx(tmp_path / "a.docx")
    run_edit(f, "--replace", "工期=质保期", "--no-backup")
    assert not (tmp_path / "a.bak.docx").exists()


def test_out_writes_elsewhere(tmp_path):
    f = make_docx(tmp_path / "a.docx")
    dst = tmp_path / "b.docx"
    run_edit(f, "--replace", "工期=质保期", "--out", str(dst))
    assert dst.exists()
    assert Document(f).tables[0].cell(1, 0).text == "工期", "原文件不该被改"


def test_list_is_read_only(tmp_path):
    f = make_docx(tmp_path / "a.docx")
    before = (tmp_path / "a.docx").stat().st_mtime_ns
    out = run_edit(f, "--list")
    assert "tables: 1" in out and "投标有效期" in out
    assert (tmp_path / "a.docx").stat().st_mtime_ns == before, "--list 不能写盘"


def test_no_operation_does_not_write(tmp_path):
    f = make_docx(tmp_path / "a.docx")
    out = run_edit(f)
    assert "没有任何改动" in out
    assert not (tmp_path / "a.bak.docx").exists()


def test_missing_file_exits(tmp_path):
    run_edit(tmp_path / "nope.docx", "--list", expect=1)


# ---------------------------------------------------------------- 批量加行 / 填表


def _make_styled_table(path, rows=2, cols=2):
    """造一张带格式的表：最后一行加粗 9pt + 底纹，供格式继承断言。"""
    doc = Document()
    t = doc.add_table(rows=rows, cols=cols)
    for ci in range(cols):
        run = t.cell(rows - 1, ci).paragraphs[0].add_run("旧值%d" % ci)
        run.bold = True
        run.font.size = Pt(9)
        tcpr = t.cell(rows - 1, ci)._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "EDEDED")
        tcpr.append(shd)
    doc.save(str(path))
    return str(path)


def test_add_rows_copies_row_format(tmp_path):
    """--add-rows 必须继承模板行的字体/字号/底纹，且文字清空。"""
    f = _make_styled_table(tmp_path / "a.docx")
    run_edit(f, "--add-rows", "0", "2", "--template-row", "1")

    t = Document(f).tables[0]
    assert len(t.rows) == 4
    for ri in (2, 3):
        for ci in range(2):
            c = t.cell(ri, ci)
            assert c.text == "", "复制出来的应是空白行"
            runs = c.paragraphs[0].runs
            assert runs and runs[0].bold is True, "加粗格式必须继承"
            assert runs[0].font.size == Pt(9), "字号必须继承"
            shd = c._tc.find(qn("w:tcPr")).find(qn("w:shd"))
            assert shd is not None and shd.get(qn("w:fill")) == "EDEDED", "底纹必须继承"


def test_add_row_plain_still_works(tmp_path):
    """旧 --add-row 行为不变：新行写入指定值。"""
    f = _make_styled_table(tmp_path / "a.docx")
    run_edit(f, "--add-row", "0", "甲", "乙")
    t = Document(f).tables[0]
    assert [t.cell(2, i).text for i in range(2)] == ["甲", "乙"]


def test_fill_batch_json(tmp_path):
    """--fill JSON：按坐标逐格写入，null 跳过不清空现值。"""
    f = _make_styled_table(tmp_path / "a.docx", rows=3, cols=3)
    doc = Document(f)
    doc.tables[0].cell(1, 2).paragraphs[0].add_run("保留值")
    doc.save(f)

    spec = {"table": 0, "start_row": 1, "start_col": 1, "values": [["甲", None], ["乙", "丙"]]}
    fp = tmp_path / "fill.json"
    fp.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    out = run_edit(f, "--fill", str(fp))
    assert "写入 3 格" in out

    t = Document(f).tables[0]
    assert t.cell(1, 1).text == "甲"
    assert t.cell(1, 2).text == "保留值", "null 不该清空现值"
    assert t.cell(2, 1).text == "乙" and t.cell(2, 2).text == "丙"


def test_fill_csv(tmp_path):
    """--fill CSV：空单元格视为跳过；默认填表 0 从 (0,0) 起。"""
    f = _make_styled_table(tmp_path / "a.docx", rows=2, cols=2)
    fp = tmp_path / "data.csv"
    fp.write_text("x,y\n,\n", encoding="utf-8")
    run_edit(f, "--fill", str(fp))

    t = Document(f).tables[0]
    assert t.cell(0, 0).text == "x" and t.cell(0, 1).text == "y"
    assert t.cell(1, 0).text == "旧值0", "CSV 空单元格不能覆盖"


def test_fill_out_of_range_warns(tmp_path):
    """越界不崩，只警告并跳过。"""
    f = _make_styled_table(tmp_path / "a.docx", rows=2, cols=2)
    spec = {"table": 0, "start_row": 1, "values": [["a", "b"], ["c", "d"]]}
    fp = tmp_path / "fill.json"
    fp.write_text(json.dumps(spec), encoding="utf-8")
    out = run_edit(f, "--fill", str(fp), expect=0)
    assert "越界" in out
    assert Document(f).tables[0].cell(1, 0).text == "a"


def test_empty_cell_borrows_format(tmp_path):
    """往空白格写值时，必须借同表已有 run 的 rPr，不能掉回默认字体。"""
    doc = Document()
    t = doc.add_table(rows=2, cols=2)
    r = t.cell(0, 0).paragraphs[0].add_run("已有值")
    r.bold = True
    f = str(tmp_path / "a.docx")
    doc.save(f)

    run_edit(f, "--cell", "0", "1", "1", "新填的值")
    c = Document(f).tables[0].cell(1, 1)
    assert c.text == "新填的值"
    runs = c.paragraphs[0].runs
    assert runs and runs[0].bold is True, "空白格应继承表内现有格式"


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
