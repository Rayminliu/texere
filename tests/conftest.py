"""测试安全网：共享的内存 docx/pdf 构造 fixture + _NoopRenderer。

把此前散落在各测试文件里重复的「内存构造 docx」helper 收敛为单一来源，并提供
一个不启动任何 Office 的 `_NoopRenderer`——它的 `render()` 恒返回 `ok=False`、
`available()` 恒返回 `(False, ...)`，于是 `generate_evidence_package` 里所有基于
导出 PDF 的派生检查（page_numbering / blank_pages / visual_drift / renderer_acceptance）
自动 SKIP，让 L2 逻辑测试能在零 Word 下跑通整条 evidence 组装链。

本文件只加测试基础设施，不改动任何生产代码。
"""

import os
import sys

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(
    0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "scripts")
)

from renderers import RenderResult  # noqa: E402  复用真实渲染器的结果契约，不重造


class _NoopRenderer:
    """测试注入渲染器：不调用任何 Office，永远导出失败。

    契约与 `renderers.Renderer` 对齐（`render(docx, pdf, timeout)` + 类方法
    `available()`），但 `render()` 直接返回 `ok=False` 的结构化结果。validate 用
    `available()` 判渲染器在不在——不在则 `renderer_acceptance` 只能 SKIP（不许冒充
    验收失败），且 `shared_pdf` 为 None，页码 / 空白页 / 视觉漂移一并 SKIP。
    """

    @classmethod
    def available(cls):
        return (False, "noop 渲染器：不启动 Office，PDF 派生检查据此 SKIP")

    def render(self, docx_path, pdf_path, timeout=300):
        return RenderResult(
            ok=False,
            pdf=None,
            renderer_name=type(self).__name__,
            renderer_version="0",
            engine_path="",
            warnings=[],
            errors=["noop renderer：不产出 PDF"],
        )


@pytest.fixture
def noop_renderer():
    """注入一个 _NoopRenderer 实例给 generate_evidence_package(renderer=...)。"""
    return _NoopRenderer()


@pytest.fixture
def make_docx(tmp_path):
    """把内存 Document 落盘并返回路径；可传若干段落文本。

    收敛各测试文件重复的 `_save_docx` / `_docx_with_text` 写法。
    """

    def _make(name="a.docx", paragraphs=()):
        doc = Document()
        for text in paragraphs:
            doc.add_paragraph(text)
        path = tmp_path / name
        doc.save(str(path))
        return str(path)

    return _make


@pytest.fixture
def make_docx_with_toc(tmp_path):
    """按 post.py 注入目录域的写法造一个含 TOC 域的 docx，返回路径。

    收敛 test_validate_units 的 `_docx_with_toc`（fldChar begin / instrText / end 三段）。
    """

    def _make(
        name="toc.docx",
        instr='TOC \\o "1-2" \\h \\z \\u',
    ):
        doc = Document()
        para = doc.add_paragraph()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        para.add_run()._r.append(begin)

        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        para.add_run()._r.append(it)

        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        para.add_run()._r.append(end)

        path = tmp_path / name
        doc.save(str(path))
        return str(path)

    return _make
