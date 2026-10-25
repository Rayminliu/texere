"""Validate CLI tests — 9-item validation coverage.

pytest -q tests/test_validate.py
"""

import hashlib
import importlib.util
import json
import os
import subprocess
import sys

import pytest

KIT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

_RENDERER_STATE = None


def _require_renderer():
    """e2e 先经 render.py --pdf 造文档：默认渲染器不可用就整组跳过（托管 CI 无 Office）。"""
    global _RENDERER_STATE
    if _RENDERER_STATE is None:
        path = os.path.join(KIT, "scripts", "renderers.py")
        spec = importlib.util.spec_from_file_location("renderers", path)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        _RENDERER_STATE = mod.WordRenderer.available()
    if not _RENDERER_STATE[0]:
        pytest.skip("渲染器不可用（%s）：e2e 需 render.py --pdf" % _RENDERER_STATE[1])


if "_sample_render" in sys.modules:  # test_patch 已加载过：复用同一实例，缓存跨文件生效
    _sample_render = sys.modules["_sample_render"]
else:
    _spec = importlib.util.spec_from_file_location(
        "_sample_render", os.path.join(KIT, "tests", "_sample_render.py")
    )
    _sample_render = importlib.util.module_from_spec(_spec)
    sys.modules["_sample_render"] = _sample_render
    _spec.loader.exec_module(_sample_render)


def _build_sample_docx(tmp_path):
    """Build a sample docx for testing.

    渲染产物进程级共享（见 tests/_sample_render.py）：全量只渲染一次，
    每条用例拿拷贝——改写型用例（--apply）不污染共享源。
    """
    return _sample_render.rendered_copy(tmp_path)


class TestPackageIntegrity:
    """Test package integrity checks."""

    @pytest.mark.word
    def test_valid_package(self):
        """Valid DOCX should pass（0.9.3 起读会话级共享 evidence，不再每用例起 Word）。"""
        report, stdout, _dir = _sample_render.shared_evidence()

        assert report["summary"]["failed"] == 0
        assert "package_integrity" in stdout
        assert report["checks"]["package_integrity"]["status"] == "PASS"

    def test_corrupted_package_exits(self, tmp_path):
        """Corrupted ZIP should fail."""
        # Create a fake ZIP file
        bad_docx = tmp_path / "bad.docx"
        bad_docx.write_bytes(b"PK\x03\x04invalid")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "validate.py"),
                str(bad_docx),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
        )

        assert result.returncode != 0
        assert "FAIL" in result.stdout or "error" in result.stderr.lower()

    def test_failure_still_reports_evidence_location(self, tmp_path):
        """T1.3 回归：验证失败时更要告诉用户证据包在哪。"""
        # 修前 print_report 在 failed>0 时自己 sys.exit(1)，main() 里排在它后面的
        # 「Evidence package saved to ...」永远执行不到——恰是最需要它的时刻。
        bad_docx = tmp_path / "bad.docx"
        bad_docx.write_bytes(b"PK\x03\x04invalid")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "validate.py"),
                str(bad_docx),
                "--out",
                str(tmp_path / "ev"),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
        )

        assert result.returncode != 0, "exit 语义不变：失败仍退出 1"
        assert "Validation FAILED" in result.stdout
        assert "Evidence package saved to" in result.stdout


@pytest.mark.word
class TestImageEmbedding:
    """Test image embedding checks."""

    def test_images_embedded(self):
        """Sample has no images but should report correctly（共享 evidence，不重复起 Word）。"""
        _report, stdout, _dir = _sample_render.shared_evidence()
        assert "image_embedding" in stdout
        assert "PASS" in stdout


@pytest.mark.word
class TestSectionCount:
    """Test section count validation."""

    def test_reasonable_sections(self):
        """Sample document should have reasonable sections（共享 evidence）。"""
        _report, stdout, _dir = _sample_render.shared_evidence()
        assert "section_count" in stdout
        assert "PASS" in stdout


@pytest.mark.word
class TestTOCField:
    """Test TOC field validation."""

    def test_toc_exists(self):
        """Sample should have TOC（共享 evidence）。"""
        _report, stdout, _dir = _sample_render.shared_evidence()

        # Should either pass (TOC exists) or skip (no TOC detected)
        assert "toc_field" in stdout


@pytest.mark.word
class TestPageNumbering:
    """Test page numbering validation."""

    def test_continuous_pages(self):
        """Pages should be continuous（共享 evidence）。"""
        _report, stdout, _dir = _sample_render.shared_evidence()
        assert "page_numbering" in stdout
        assert "PASS" in stdout


@pytest.mark.word
class TestBlankPages:
    """Test blank page detection."""

    def test_no_blank_pages(self):
        """Sample should have no blank pages（共享 evidence）。"""
        _report, stdout, _dir = _sample_render.shared_evidence()
        assert "blank_pages" in stdout
        # Sample has 4 pages, should pass with max_empty=0
        assert "PASS" in stdout

    def test_max_empty_relaxes_threshold(self, tmp_path):
        """--max-empty should relax the threshold（不同输入：保持独立 spawn）。"""
        docx = _build_sample_docx(tmp_path)

        # This should work even if there are blank pages
        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "validate.py"),
                docx,
                "--max-empty",
                "5",
                "--quiet",
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
        )

        assert result.returncode == 0


@pytest.mark.word
class TestWordAcceptance:
    """Test Word acceptance validation."""

    def test_word_opens_and_exports(self):
        """Document should open in Word and export PDF（共享 evidence）。"""
        _report, stdout, _dir = _sample_render.shared_evidence()
        assert "renderer_acceptance" in stdout
        assert "PASS" in stdout


@pytest.mark.word
class TestVisualDrift:
    """Test visual drift detection."""

    def test_drift_skipped_without_baseline(self):
        """Should skip drift check without baseline（共享 evidence）。"""
        _report, stdout, _dir = _sample_render.shared_evidence()
        assert "visual_drift" in stdout
        # Should PASS with skip message
        assert "PASS" in stdout or "跳过" in stdout

    def test_drift_with_profile(self, tmp_path):
        """Should use profile if provided（不同输入：保持独立 spawn）。"""
        docx = _build_sample_docx(tmp_path)
        profile = os.path.join(KIT, "profiles", "formal-cn-v1.json")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "validate.py"),
                docx,
                "--profile",
                profile,
                "--quiet",
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
        )

        assert result.returncode == 0


@pytest.mark.word
class TestEvidencePackage:
    """Test evidence package generation.

    四件套断言（目录 / report.json / signature / 截图）合并到一次 validate
    运行：同一份产物不该拆成四次 Word 往返。
    """

    def test_evidence_artifacts_created(self):
        """Evidence directory, report.json, signature and screenshots all created（共享 evidence）。"""
        _report, _stdout, evidence_dir = _sample_render.shared_evidence()

        assert os.path.isdir(evidence_dir)
        assert any(os.scandir(evidence_dir))

        report_path = os.path.join(evidence_dir, "report.json")
        assert os.path.exists(report_path)

        # Verify JSON structure
        report = _report
        assert "metadata" in report
        assert "checks" in report
        assert "summary" in report
        assert "total" in report["summary"]
        assert "passed" in report["summary"]
        assert "failed" in report["summary"]

        sig_path = os.path.join(evidence_dir, "signature")
        assert os.path.exists(sig_path)

        # Verify signature content
        with open(sig_path, encoding="utf-8") as f:
            content = f.read()
        assert "document_hash" in content
        assert "checks_passed" in content
        # report.json 的 hash 也要进证据清单：否则改报告结论不会破坏证据
        assert "report_hash" in content
        assert "cryptographic" in content  # 明说是 checksum manifest，不是签名

        # 全量文件清单：报告与全部截图都必须入清单（此前只盖 docx 与 report，截图是漏项）
        assert "tool_version:" in content
        assert "generated_at:" in content
        assert "file[report.json]:" in content
        for name in sorted(os.listdir(evidence_dir)):
            if name.startswith("page-"):
                assert "file[%s]:" % name in content, "截图 %s 未入证据清单" % name

        # Should have at least one screenshot
        png_files = [
            n for n in os.listdir(evidence_dir) if n.startswith("page-") and n.endswith(".png")
        ]
        assert len(png_files) >= 1


class TestExitCodes:
    """Test exit code behavior."""

    @pytest.mark.word
    def test_success_exit_code(self):
        """Successful validation should exit 0（共享 evidence：check=True 已证 exit 0）。"""
        _report, _stdout, _dir = _sample_render.shared_evidence()
        # shared_evidence 内部 subprocess.run(check=True)：能拿到返回值即退出码 0
        assert _report["summary"]["failed"] == 0

    def test_failure_exit_code(self, tmp_path):
        """Failed validation should exit 1."""
        # Use invalid file
        bad_file = tmp_path / "notexist.docx"

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "validate.py"),
                str(bad_file),
            ],
            capture_output=True,
        )

        assert result.returncode != 0


class TestHelpAndVersion:
    """Test CLI help and version."""

    def test_help_flag(self):
        """--help should show usage."""
        result = subprocess.run(
            [sys.executable, os.path.join(KIT, "scripts", "validate.py"), "--help"],
            capture_output=True,
            text=True,
            encoding="utf-8",
        )

        assert result.returncode == 0
        assert "usage:" in result.stdout.lower()
        assert "validate.py" in result.stdout

    def test_positional_argument_required(self, tmp_path):
        """Missing positional argument should exit with error."""
        result = subprocess.run(
            [sys.executable, os.path.join(KIT, "scripts", "validate.py")],
            capture_output=True,
            text=True,
            encoding="utf-8",
        )

        assert result.returncode != 0


@pytest.mark.word
class TestProfileValidation:
    """Test profile integration（0.9.3：内置 profile 合并为一次参数化运行）。

    旧版每个 profile 各 spawn 一次 validate = 多次 Word 往返；profile 的差异只
    影响 --profile 文件本身（metadata 附带与 baseline 回退），同一份文档上行为
    仅在契约字段上有别——参数化共享 docx、逐 profile 断言 exit 0 语义不变。
    """

    @pytest.mark.parametrize(
        "profile_name",
        ["formal-cn-v1.json", "gongwen-v1.json", "application-v1.json"],
    )
    def test_profile_works(self, tmp_path, profile_name):
        docx = _build_sample_docx(tmp_path)
        profile = os.path.join(KIT, "profiles", profile_name)

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "validate.py"),
                docx,
                "--profile",
                profile,
                "--quiet",
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
        )

        assert result.returncode == 0, result.stdout + result.stderr


@pytest.mark.word
class TestReportStructure:
    """Test report JSON structure.

    三项断言读的是同一份 report.json，合并到一次 validate 运行。
    """

    def test_report_structure_complete(self):
        """Report contains all 9 checks, a timestamp and the tool version（共享 evidence）。"""
        report, _stdout, _dir = _sample_render.shared_evidence()

        expected_checks = [
            "package_integrity",
            "source_content",
            "image_embedding",
            "section_count",
            "toc_field",
            "page_numbering",
            "blank_pages",
            "renderer_acceptance",
            "visual_drift",
        ]

        for check_name in expected_checks:
            assert check_name in report["checks"], f"Missing check: {check_name}"

        assert "timestamp" in report["metadata"]
        # Should be ISO format
        assert "T" in report["metadata"]["timestamp"]

        assert "tool_version" in report["metadata"]
        # 版本从 scripts/_version.py 单一来源读取，发版升号时不必再改本测试
        import re

        with open(os.path.join(KIT, "scripts", "_version.py"), encoding="utf-8") as f:
            m = re.search(r'__version__\s*=\s*"([^"]+)"', f.read())
        assert m, "scripts/_version.py 里找不到 __version__"
        assert report["metadata"]["tool_version"] == m.group(1)


def _sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def _make_docx(tmp_path, name="in.docx"):
    """最小可 Word 打开的 docx（含一个 TOC 域，便于域刷新改变字节）。"""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p = tmp_path / name
    doc = Document()
    doc.add_paragraph("Texere read-only acceptance regression")
    para = doc.add_paragraph()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    para.add_run()._r.append(begin)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-2" \\h \\z \\u'
    para.add_run()._r.append(it)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    para.add_run()._r.append(end)
    doc.save(str(p))
    return str(p)


@pytest.mark.word
class TestFinalizeIsReadOnly:
    """钉死 read-only 验收契约：默认不改输入，--save-updated-fields 才写回。

    否则将来某维护者在 finalize.py 里随手加回 doc.Save()，immutable acceptance
    这条承诺就会静默失效。"""

    def test_default_does_not_mutate_input(self, tmp_path):
        # finalize 默认要出 PDF（Word 渲染）；无渲染器环境（托管 CI）跳过，不算失败
        _require_renderer()
        src = _make_docx(tmp_path, "ro.docx")
        before = _sha256(src)
        r = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "finalize.py"),
                src,
                str(tmp_path / "ro.pdf"),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
        )
        assert r.returncode == 0, r.stderr
        after = _sha256(src)
        assert before == after, "finalize 默认不应修改输入 docx（验收必须 immutable）"
        assert (tmp_path / "ro.pdf").exists()

    def test_save_updated_fields_mutates_input(self, tmp_path):
        _require_renderer()
        src = _make_docx(tmp_path, "su.docx")
        before = _sha256(src)
        r = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "finalize.py"),
                src,
                str(tmp_path / "su.pdf"),
                "--save-updated-fields",
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
        )
        assert r.returncode == 0, r.stderr
        after = _sha256(src)
        assert before != after, "--save-updated-fields 应把刷新后的域写回原 docx"


@pytest.mark.word
def test_provenance_recorded(tmp_path):
    """证据必须回答「哪个契约 + 哪个转换器」（不同输入：--config/--reference 独立 spawn）。

    0.9.3：断言本体不变，但 docx 改用共享渲染产物直接拷贝——渲染只做一次
    （rendered_copy 已是进程级缓存），这里只省掉重复渲染，validate 仍独立跑。
    """
    docx = _build_sample_docx(tmp_path)
    evidence_dir = tmp_path / "evidence"
    result = subprocess.run(
        [
            sys.executable,
            os.path.join(KIT, "scripts", "validate.py"),
            docx,
            "--out",
            str(evidence_dir),
            "--quiet",
            "--config",
            os.path.join(KIT, "assets", "sample_config.json"),
            "--reference",
            os.path.join(KIT, "assets", "ref.docx"),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert result.returncode == 0, result.stderr
    report = json.load(open(evidence_dir / "report.json", encoding="utf-8"))
    md = report["metadata"]
    assert md["config_sha256"] and len(md["config_sha256"]) == 64
    assert md["reference_sha256"] and len(md["reference_sha256"]) == 64
    assert md["pandoc_version"] and md["pandoc_version"] != "unknown"
    assert md["cli"], "CLI 命令行应进 metadata"
    sig = (evidence_dir / "signature").read_text(encoding="utf-8")
    for key in ("config_sha256:", "reference_sha256:", "pandoc_version:"):
        assert key in sig, "signature 缺 %s" % key
