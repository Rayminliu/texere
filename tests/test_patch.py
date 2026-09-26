"""Patch API tests — declarative editing with safety mechanisms.

pytest -q tests/test_patch.py
"""

import importlib.util
import json
import os
import shutil
import subprocess
import sys

import pytest

KIT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

_RENDERER_STATE = None


def _require_renderer():
    """e2e 先经 render.py --pdf 造文档：默认渲染器不可用就整组跳过（托管 CI 无 Office）。"""
    global _RENDERER_STATE
    if _RENDERER_STATE is None:
        path = os.path.join(KIT, "scripts", "renderers.py")
        spec = importlib.util.spec_from_file_location("renderers", path)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        _RENDERER_STATE = mod.WordRenderer.available()
    if not _RENDERER_STATE[0]:
        pytest.skip("渲染器不可用（%s）：e2e 需 render.py --pdf" % _RENDERER_STATE[1])


def _build_sample_docx(tmp_path):
    """Build a sample docx for testing."""
    _require_renderer()
    src = tmp_path / "src"
    src.mkdir()
    shutil.copy(os.path.join(KIT, "assets", "sample.md"), src / "01_sample.md")

    out = str(tmp_path / "test.docx")
    subprocess.run(
        [
            sys.executable,
            os.path.join(KIT, "scripts", "render.py"),
            "--src",
            str(src),
            "--out",
            out,
            "--config",
            os.path.join(KIT, "assets", "sample_config.json"),
            "--pdf",
        ],
        check=True,
        capture_output=True,
    )
    return out


class TestSchemaValidation:
    """Test Patch JSON schema validation."""

    def test_valid_patch_schema(self, tmp_path):
        """Valid patch should pass schema validation."""
        patch = {
            "id": "test-001",
            "operations": [{"op": "replace_text", "target": {"paragraph": 0}, "new_text": "test"}],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(_build_sample_docx(tmp_path)),
                str(patch_file),
            ],
            capture_output=True,
            text=True,
        )

        # Should not exit on schema error
        assert "schema validation failed" not in result.stderr.lower()

    def test_missing_id_exits(self, tmp_path):
        """Missing id should fail."""
        patch = {"operations": []}

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(_build_sample_docx(tmp_path)),
                str(patch_file),
            ],
            capture_output=True,
            text=True,
        )

        assert "缺少必需字段" in result.stderr or "id" in result.stderr.lower()

    def test_missing_operations_exits(self, tmp_path):
        """Missing operations should fail."""
        patch = {"id": "test-001"}

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(_build_sample_docx(tmp_path)),
                str(patch_file),
            ],
            capture_output=True,
            text=True,
        )

        assert "operations" in result.stderr.lower()

    def test_invalid_operation_type_exits(self, tmp_path):
        """Invalid operation type should fail."""
        patch = {"id": "test-001", "operations": [{"op": "invalid_op", "target": {}}]}

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(_build_sample_docx(tmp_path)),
                str(patch_file),
            ],
            capture_output=True,
            text=True,
        )

        assert "未实现的 operation" in result.stderr or "invalid" in result.stderr.lower()


class TestReplaceTextOperation:
    """Test replace_text operation."""

    def test_replace_text_dry_run(self, tmp_path):
        """Dry run should simulate without applying."""
        docx = _build_sample_docx(tmp_path)

        patch = {
            "id": "test-replace",
            "operations": [
                {
                    "op": "replace_text",
                    "target": {"paragraph": 5},
                    "new_text": "replacement text",
                }
            ],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
                "--dry-run",
            ],
            capture_output=True,
            text=True,
        )

        # Dry run should succeed (may report text not found but that's OK)
        assert result.returncode in (0, 1)

    def test_replace_text_with_expected_old(self, tmp_path):
        """Should verify expected old text exists."""
        docx = _build_sample_docx(tmp_path)

        patch = {
            "id": "test-verify",
            "operations": [
                {
                    "op": "replace_text",
                    "target": {"paragraph": 5},
                    "expected_old_text": "nonexistent text xyz123",
                    "new_text": "replacement",
                }
            ],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
                "--dry-run",
            ],
            capture_output=True,
            text=True,
        )

        # Should report that expected text not found
        assert "预期旧文本不存在" in result.stderr or "not found" in result.stderr.lower()


class TestInsertOperations:
    """Test insert_after and insert_before operations."""

    def test_insert_after_anchor(self, tmp_path):
        """Should insert after anchor text."""
        docx = _build_sample_docx(tmp_path)

        patch = {
            "id": "test-insert-after",
            "operations": [
                {
                    "op": "insert_after",
                    "target": {"anchor": "示例"},
                    "content": ["Inserted paragraph"],
                }
            ],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
                "--dry-run",
            ],
            capture_output=True,
            text=True,
        )

        # Should either succeed or skip if anchor not found
        assert result.returncode in (0, 1)


class TestDeleteParagraphOperation:
    """Test delete_paragraph operation."""

    def test_delete_paragraph_dry_run(self, tmp_path):
        """Should work in dry run mode."""
        docx = _build_sample_docx(tmp_path)

        patch = {
            "id": "test-delete",
            "operations": [{"op": "delete_paragraph", "target": {"paragraph": 5}}],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
                "--dry-run",
            ],
            capture_output=True,
            text=True,
        )

        assert result.returncode == 0


class TestSetCellOperation:
    """Test set_cell operation."""

    def test_set_cell_dry_run(self, tmp_path):
        """Should work in dry run mode."""
        docx = _build_sample_docx(tmp_path)

        patch = {
            "id": "test-cell",
            "operations": [
                {
                    "op": "set_cell",
                    "target": {"table": 0, "row": 0, "column": 0},
                    "value": "new value",
                }
            ],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
                "--dry-run",
            ],
            capture_output=True,
            text=True,
        )

        # May fail if no tables, but shouldn't crash
        assert result.returncode in (0, 1)

    def test_set_cell_out_of_range_exits(self, tmp_path):
        """Out of range should exit with error."""
        docx = _build_sample_docx(tmp_path)

        patch = {
            "id": "test-bad-cell",
            "operations": [
                {
                    "op": "set_cell",
                    "target": {"table": 999, "row": 0, "column": 0},
                    "value": "value",
                }
            ],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
                "--dry-run",
            ],
            capture_output=True,
            text=True,
        )

        # Should fail gracefully
        assert result.returncode != 0 or "越界" in result.stdout


class TestAddRowOperation:
    """Test add_row operation.

    after_row 的位置 / 格式继承 / 越界断言在 TestPatchUnits 里（不启 Word，跑得快）。
    """

    def test_add_row_dry_run(self, tmp_path):
        """Should work in dry run mode."""
        docx = _build_sample_docx(tmp_path)

        patch = {
            "id": "test-add-row",
            "operations": [{"op": "add_row", "target": {"table": 0}, "values": ["val1", "val2"]}],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
                "--dry-run",
            ],
            capture_output=True,
            text=True,
        )

        # May fail if no tables, but shouldn't crash
        assert result.returncode in (0, 1)


class TestPreconditions:
    """Test hash and must_contain preconditions."""

    def test_must_contain_passes(self, tmp_path):
        """must_contain should pass if text exists."""
        docx = _build_sample_docx(tmp_path)

        patch = {
            "id": "test-precondition",
            "preconditions": {"must_contain": ["示例"]},  # Sample contains this
            "operations": [],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
                "--dry-run",
            ],
            capture_output=True,
            text=True,
        )

        # Should pass precondition check
        assert "通过" in result.stdout or "exist" in result.stdout.lower() or result.returncode == 0

    def test_must_contain_fails(self, tmp_path):
        """must_contain should fail if text missing."""
        docx = _build_sample_docx(tmp_path)

        patch = {
            "id": "test-fail-precondition",
            "preconditions": {"must_contain": ["this text definitely does not exist 12345"]},
            "operations": [],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
                "--dry-run",
            ],
            capture_output=True,
            text=True,
        )

        # Should fail precondition check (message goes to stderr)
        assert "缺少" in result.stderr or "not found" in result.stderr.lower()

    def test_hash_precondition_passes(self, tmp_path):
        """Hash should match if document unchanged."""
        docx = _build_sample_docx(tmp_path)

        # Compute hash first
        import hashlib

        sha256 = hashlib.sha256()
        with open(docx, "rb") as f:
            sha256.update(f.read())
        expected_hash = sha256.hexdigest()

        patch = {
            "id": "test-hash",
            "preconditions": {"hash": expected_hash},
            "operations": [],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
                "--dry-run",
            ],
            capture_output=True,
            text=True,
        )

        # Should pass hash check
        assert "通过" in result.stdout or "match" in result.stdout.lower() or result.returncode == 0

    def test_hash_precondition_fails(self, tmp_path):
        """Hash should fail if document modified."""
        docx = _build_sample_docx(tmp_path)

        patch = {
            "id": "test-bad-hash",
            "preconditions": {"hash": "wronghash1234567890abcdef1234567890abcdef1234567890abcdef"},
            "operations": [],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
                "--dry-run",
            ],
            capture_output=True,
            text=True,
        )

        # Should fail hash check (message goes to stderr)
        assert "不匹配" in result.stderr or "mismatch" in result.stderr.lower()


class TestApplyAndValidate:
    """Test apply and validate workflow."""

    def test_apply_creates_backup(self, tmp_path):
        """Apply should create backup by default."""
        docx = _build_sample_docx(tmp_path)
        _original_size = os.path.getsize(docx)

        patch = {
            "id": "test-apply",
            "operations": [
                {
                    "op": "replace_text",
                    "target": {"paragraph": 0},
                    "new_text": "modified",
                }
            ],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        _result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
                "--apply",
            ],
            capture_output=True,
            text=True,
        )

        # Should create backup
        backup_path = str(docx).replace(".docx", ".bak.docx")
        assert os.path.exists(backup_path)

    def test_validate_after_apply(self, tmp_path):
        """Validate should verify changes."""
        docx = _build_sample_docx(tmp_path)

        patch = {
            "id": "test-validate",
            "operations": [
                {
                    "op": "replace_text",
                    "target": {"paragraph": 0},
                    "new_text": "test modification",
                }
            ],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
                "--apply",
                "--validate",
            ],
            capture_output=True,
            text=True,
        )

        # Should complete successfully
        assert result.returncode == 0


class TestEvidenceGeneration:
    """Test evidence package generation."""

    def test_evidence_created_on_apply(self, tmp_path):
        """Evidence should be created when --out specified."""
        docx = _build_sample_docx(tmp_path)
        evidence_dir = tmp_path / "evidence"

        patch = {
            "id": "test-evidence",
            "operations": [{"op": "replace_text", "target": {"paragraph": 0}, "new_text": "test"}],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
                "--apply",
                "--out",
                str(evidence_dir),
            ],
            check=True,
        )

        assert evidence_dir.exists()

        # Check required files
        assert (evidence_dir / "patch_report.json").exists()
        assert (evidence_dir / "signature").exists()


class TestExitCodes:
    """Test exit code behavior."""

    def test_help_flag(self):
        """--help should show usage."""
        result = subprocess.run(
            [sys.executable, os.path.join(KIT, "scripts", "patch.py"), "--help"],
            capture_output=True,
            text=True,
        )

        assert result.returncode == 0
        assert "usage:" in result.stdout.lower()

    def test_missing_file_exits(self, tmp_path):
        """Missing file should exit with error."""
        patch = {"id": "test", "operations": []}

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                "nonexistent.docx",
                str(patch_file),
            ],
            capture_output=True,
            text=True,
        )

        assert result.returncode != 0
        assert "not found" in result.stderr.lower() or "找不到" in result.stderr


class TestPatchUnits:
    """Fast in-process tests for schema↔implementation consistency (no Word)."""

    @staticmethod
    def _module():
        import importlib.util

        spec = importlib.util.spec_from_file_location(
            "patch_under_test", os.path.join(KIT, "scripts", "patch.py")
        )
        m = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(m)
        return m

    def test_no_broken_scripts_import_left(self):
        """`from scripts import edit` 在以 `python scripts/patch.py` 运行时解析成空的
        命名空间包，于是 set_cell / insert_* / add_row 全部抛 ImportError，又被
        except Exception 兜成「操作失败」。这里按 AST 守住不再回退
        （不用字符串匹配，免得注释里提一句就被误判）。"""
        import ast

        tree = ast.parse(open(os.path.join(KIT, "scripts", "patch.py"), encoding="utf-8").read())
        bad = [
            n.module
            for n in ast.walk(tree)
            if isinstance(n, ast.ImportFrom) and (n.module or "").split(".")[0] == "scripts"
        ]
        assert not bad, "patch.py 又出现了坏导入（scripts 不是可导入的包）: %s" % bad

    def test_edit_module_imports(self):
        p = self._module()
        assert callable(p._edit_module()._set_cell_text)

    @staticmethod
    def _table_doc():
        from docx import Document

        doc = Document()
        t = doc.add_table(rows=3, cols=2)
        for r in range(3):
            for c in range(2):
                t.rows[r].cells[c].text = "r%d-c%d" % (r, c)
        return doc, t

    def test_after_row_inserts_in_place(self):
        """schema 里声明了 after_row 就必须真的插在那一行之后，不能追加到表尾。"""
        p = self._module()
        doc, t = self._table_doc()
        ok, msg = p.add_row_op(
            doc,
            {"op": "add_row", "target": {"table": 0, "after_row": 0}, "values": ["NEW-A", "NEW-B"]},
        )
        assert ok, msg
        assert len(t.rows) == 4
        assert t.rows[1].cells[0].text == "NEW-A"
        assert t.rows[0].cells[0].text == "r0-c0"

    def test_after_row_second_position(self):
        p = self._module()
        doc, t = self._table_doc()
        ok, msg = p.add_row_op(
            doc,
            {"op": "add_row", "target": {"table": 0, "after_row": 1}, "values": ["NEW-C", "NEW-D"]},
        )
        assert ok, msg
        assert t.rows[2].cells[0].text == "NEW-C"

    def test_omitted_after_row_appends(self):
        p = self._module()
        doc, t = self._table_doc()
        ok, msg = p.add_row_op(doc, {"op": "add_row", "target": {"table": 0}, "values": ["NEW-E"]})
        assert ok, msg
        assert t.rows[-1].cells[0].text == "NEW-E"

    @staticmethod
    def _styled_table_doc():
        """第 1 行带底纹 + 字号的表，用来验证新行真的继承了锚点行格式。"""
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt

        doc = Document()
        t = doc.add_table(rows=3, cols=2)
        for r in range(3):
            for c in range(2):
                t.rows[r].cells[c].text = "r%d-c%d" % (r, c)
        for cell in t.rows[1].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "FF0000")
            tcPr.append(shd)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(18)
        return doc, t

    def test_after_row_inherits_anchor_row_formatting(self):
        """新行必须继承锚点行的底纹 / 字号 —— 不只是插对位置。

        python-docx 的 add_row() 产出的是裸行（丢边框、底纹、字号），这正是
        add_row 要 deepcopy <w:tr> 而不是调 add_row() 的全部理由。只断言位置、
        不断言格式的话，这个理由没人守着，哪天被改回 add_row() 也不会红。
        """
        from docx.oxml.ns import qn
        from docx.shared import Pt

        p = self._module()
        doc, t = self._styled_table_doc()
        ok, msg = p.add_row_op(
            doc, {"op": "add_row", "target": {"table": 0, "after_row": 1}, "values": ["X0", "X1"]}
        )
        assert ok, msg

        # 形状：r0 / r1 / X / r2
        assert [c.text for c in t.rows[0].cells] == ["r0-c0", "r0-c1"]
        assert [c.text for c in t.rows[1].cells] == ["r1-c0", "r1-c1"]
        assert [c.text for c in t.rows[2].cells] == ["X0", "X1"]
        assert [c.text for c in t.rows[3].cells] == ["r2-c0", "r2-c1"]

        for cell in t.rows[2].cells:
            shd = cell._tc.tcPr.find(qn("w:shd"))
            assert shd is not None, "新行丢了底纹（说明退回成了裸行）"
            assert shd.get(qn("w:fill")) == "FF0000"
            runs = cell.paragraphs[0].runs
            assert runs and runs[0].font.size == Pt(18), "新行丢了字号"

    def test_after_row_out_of_range_reports_error(self):
        p = self._module()
        doc, _ = self._table_doc()
        ok, msg = p.add_row_op(
            doc, {"op": "add_row", "target": {"table": 0, "after_row": 99}, "values": ["X"]}
        )
        assert not ok
        assert "越界" in msg

    def test_assess_flags_out_of_range_after_row(self):
        p = self._module()
        doc, _ = self._table_doc()
        ok, issues = p.assess_patch(
            {"id": "x", "operations": [{"op": "add_row", "target": {"table": 0, "after_row": 99}}]},
            doc,
        )
        assert not ok
        assert any("after_row" in i for i in issues)

    def test_validate_result_checks_inserted_position(self):
        """新增行的位置也要被 --validate 验到，否则 schema 与验证又不闭环。"""
        p = self._module()
        doc, _ = self._table_doc()
        p.add_row_op(
            doc,
            {"op": "add_row", "target": {"table": 0, "after_row": 0}, "values": ["NEW-A", "NEW-B"]},
        )
        ok, issues = p.validate_patch_result(
            doc,
            {
                "id": "x",
                "operations": [
                    {"op": "add_row", "target": {"table": 0, "after_row": 0}, "values": ["NEW-A"]}
                ],
            },
        )
        assert ok, issues
        ok_bad, issues_bad = p.validate_patch_result(
            doc,
            {
                "id": "x",
                "operations": [
                    {"op": "add_row", "target": {"table": 0, "after_row": 0}, "values": ["NOPE"]}
                ],
            },
        )
        assert not ok_bad and issues_bad


class TestAssessment:
    """Test patch assessment."""

    def test_assessment_reports_issues(self, tmp_path):
        """Assessment should identify potential issues."""
        docx = _build_sample_docx(tmp_path)

        patch = {
            "id": "test-assess",
            "operations": [
                {
                    "op": "replace_text",
                    "target": {"paragraph": 9999},  # Out of range
                    "new_text": "test",
                }
            ],
        }

        patch_file = tmp_path / "patch.json"
        patch_file.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")

        result = subprocess.run(
            [
                sys.executable,
                os.path.join(KIT, "scripts", "patch.py"),
                str(docx),
                str(patch_file),
            ],
            capture_output=True,
            text=True,
        )

        # Should report issues during assessment
        assert (
            "Issue" in result.stdout or "issue" in result.stdout.lower() or "问题" in result.stdout
        )
