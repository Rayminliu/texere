"""后处理冒烟断言：只需 pandoc + python-docx，不需要 Word / PDF，秒级可跑。

  pytest -q

覆盖 README「排版六条军规」里能被机器判定的部分：
跨页重复表头、题注居中、封面与目录注入、分节页码。
"""

import json
import os
import shutil
import subprocess
import sys

import pytest

KIT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LUA_FILTER = os.path.join(KIT, "scripts", "filters", "captions.lua")

pytestmark = pytest.mark.skipif(shutil.which("pandoc") is None, reason="需要 pandoc")


def _pandoc_cmd(md, body, res, cfg=None):
    """与 render.py 保持一致：带上 lua filter，并在配置题注关键字时传 -M。"""
    cmd = [
        "pandoc",
        md,
        "-o",
        body,
        "--reference-doc=" + os.path.join(KIT, "assets", "ref.docx"),
        "--resource-path=" + res,
        "-f",
        "markdown+pipe_tables+raw_html",
        "--wrap=none",
    ]
    if os.path.exists(LUA_FILTER):
        cmd.append("--lua-filter=" + LUA_FILTER)
        cw = (cfg or {}).get("caption_words") or {}
        for key, meta_name in (
            ("table", "dk-table-words"),
            ("figure", "dk-figure-words"),
        ):
            if cw.get(key):
                cmd += ["-M", "%s=%s" % (meta_name, ",".join(cw[key]))]
    return cmd


def _build(tmp_path):
    src = tmp_path / "src"
    src.mkdir()
    shutil.copy(os.path.join(KIT, "assets", "sample.md"), src / "01_sample.md")
    body = str(tmp_path / "body.docx")
    out = str(tmp_path / "out.docx")
    subprocess.run(
        _pandoc_cmd(str(src / "01_sample.md"), body, str(src)),
        check=True,
        capture_output=True,
    )
    subprocess.run(
        [
            sys.executable,
            os.path.join(KIT, "scripts", "post.py"),
            body,
            out,
            os.path.join(KIT, "assets", "sample_config.json"),
        ],
        check=True,
        capture_output=True,
    )
    return out


def _build_md(tmp_path, md_text, cfg):
    """用自定义 md 与 config 走一遍 pandoc + post.py，返回输出 docx 路径。"""
    src = tmp_path / "src"
    src.mkdir(exist_ok=True)
    (src / "01.md").write_text(md_text, encoding="utf-8")
    cfg_path = tmp_path / "cfg.json"
    cfg_path.write_text(json.dumps(cfg, ensure_ascii=False), encoding="utf-8")
    body, out = str(tmp_path / "body.docx"), str(tmp_path / "out.docx")
    subprocess.run(
        _pandoc_cmd(str(src / "01.md"), body, str(src), cfg),
        check=True,
        capture_output=True,
    )
    subprocess.run(
        [
            sys.executable,
            os.path.join(KIT, "scripts", "post.py"),
            body,
            out,
            str(cfg_path),
        ],
        check=True,
        capture_output=True,
    )
    return out


@pytest.fixture(scope="module")
def docx_path(tmp_path_factory):
    return _build(tmp_path_factory.mktemp("build"))


def test_tables_repeat_header(docx_path):
    """军规 3：每个表格首行必须带 w:tblHeader，跨页才会重复表头。"""
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(docx_path)
    assert doc.tables, "样例里应当有表格"
    for i, tbl in enumerate(doc.tables):
        trPr = tbl.rows[0]._tr.find(qn("w:trPr"))
        assert trPr is not None, "表 %d 首行缺少 trPr" % i
        assert trPr.find(qn("w:tblHeader")) is not None, "表 %d 首行缺少 tblHeader" % i


def test_captions_centered(docx_path):
    """军规 4：「表 x-y 标题」应被识别为表题并居中（sample.md 里有 2 处）。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document(docx_path)
    caps = [
        p
        for p in doc.paragraphs
        if p.text.strip().startswith("表 ")
        and "响应表" in p.text
        or p.text.strip().startswith("表 2-1")
    ]
    assert len(caps) == 2, "应识别到 2 条表题，实际 %d" % len(caps)
    for p in caps:
        assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER, "表题未居中: %s" % p.text


def test_english_caption_styled(tmp_path):
    """英文题注（Table 1: ...）同样被识别并居中——neutral-en 语境下的管线事实。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    md = (
        "Intro paragraph before the table.\n"
        "\n"
        "Table 1: Quarterly milestone summary\n"
        "\n"
        "| Milestone | Status |\n"
        "|-----------|--------|\n"
        "| M1        | Done   |\n"
    )
    docx = _build_md(tmp_path, md, {})
    doc = Document(docx)
    caps = [p for p in doc.paragraphs if p.text.strip().startswith("Table 1:")]
    assert len(caps) == 1, "英文题注应被识别，实际 %d" % len(caps)
    assert caps[0].alignment == WD_ALIGN_PARAGRAPH.CENTER, "英文表题未居中: %s" % caps[0].text


def test_cover_and_toc_injected(docx_path):
    """封面行与目录域必须注入，且目录在正文之前。"""
    from docx import Document

    doc = Document(docx_path)
    texts = [p.text.strip() for p in doc.paragraphs]
    assert "投 标 文 件" in texts, "封面标题未注入"
    toc_idx = next((i for i, t in enumerate(texts) if "目录" in t), None)
    assert toc_idx is not None, "目录标题未注入"
    body_idx = next((i for i, t in enumerate(texts) if "投标函" in t), None)
    assert body_idx is not None, "未找到正文第一章"
    assert toc_idx < body_idx, "目录必须排在正文之前"


def _has_toc_field(path):
    """document.xml 里是否有 TOC 域（instrText 含 'TOC \\o'）。"""
    import zipfile

    with zipfile.ZipFile(path) as z:
        return "TOC \\o" in z.read("word/document.xml").decode("utf-8")


MD_WITH_HEADINGS = """# 第一章 总则

正文一二三。

## 1.1 细则

细则内容。
"""


def test_toc_false_skips_toc_keeps_headings(tmp_path):
    """config "toc": false：不插目录域、不分节，但 Heading 样式照常保留。

    这是通知/公示类短文档的正路；旧写法只能不写 # 绕过，代价是丢标题样式。
    """
    from docx import Document

    out = _build_md(tmp_path, MD_WITH_HEADINGS, {"toc": False})
    assert not _has_toc_field(out), "toc:false 时不该有 TOC 域"

    doc = Document(out)
    styles = [p.style.name for p in doc.paragraphs if p.text.strip()]
    assert "Heading 1" in styles, "toc:false 不能把标题降级成普通段落"
    assert "Heading 2" in styles
    assert len(doc.sections) == 1, "无封面又无目录时不该分节（会多出空白首页）"


def test_toc_false_with_cover_still_sections(tmp_path):
    """toc:false + 封面：封面单独成节，页码体系照常，只是没有目录页。"""
    from docx import Document

    cfg = {"toc": False, "cover": [["CoverTitle", "通 知"]]}
    out = _build_md(tmp_path, MD_WITH_HEADINGS, cfg)
    assert not _has_toc_field(out)
    doc = Document(out)
    assert len(doc.sections) == 2, "有封面时仍应分节（封面/正文各自页码）"
    assert "通 知" in [p.text.strip() for p in doc.paragraphs]


def test_never_touches_text(tmp_path):
    """核心契约：post.py 只改版式，一个字都不改内容。

    曾经开启过自动编号功能，它改写过题注文字（`附件 8-1 …` → `图 8-1 附件 8-1 …`）
    并篡改过正文（`**表层…**` → `表 3-5 层…`）。该功能已删除，这里守住契约。
    """
    from docx import Document

    md = (
        "# 第一章\n\n表 9-9 商务条款响应表\n\n| a |\n|:--|\n| 1 |\n\n"
        "- **表层（边缘轻算力）：** 基于公开预训练模型迁移微调。\n"
    )
    out = _build_md(tmp_path, md, {})
    texts = [p.text.strip() for p in Document(out).paragraphs]
    assert "表 9-9 商务条款响应表" in texts, "题注文字被改动了"
    assert any(t.startswith("表层（边缘轻算力）") for t in texts), "正文被改动了：%s" % [
        t for t in texts if "表层" in t
    ]


def test_two_sections_with_page_number(docx_path):
    """封面/目录为第 1 节，正文为第 2 节，正文节页脚有页码域。"""
    from docx import Document

    doc = Document(docx_path)
    assert len(doc.sections) == 2, "应为 2 节，实际 %d" % len(doc.sections)
    footer_xml = doc.sections[-1].footer.paragraphs[0]._p.xml
    assert "PAGE" in footer_xml, "正文节页脚缺少页码域"


def test_caption_variants_centered(tmp_path):
    """全角编号、英文关键字都要识别为题注（居中），但文字保持原样。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    md = (
        "# 第一章 测试\n\n"
        "表1.1 窄格式表题\n\n| a |\n|:--|\n| 1 |\n\n"
        "表 １－１ 全角编号\n\n| a |\n|:--|\n| 1 |\n\n"
        "Figure 1-1 英文图注\n\n"
        "Table 1-9 另一张表\n\n| a |\n|:--|\n| 1 |\n"
    )
    out = _build_md(tmp_path, md, {})
    doc = Document(out)
    texts = [p.text.strip() for p in doc.paragraphs]
    for want in (
        "表1.1 窄格式表题",
        "表 １－１ 全角编号",
        "Figure 1-1 英文图注",
        "Table 1-9 另一张表",
    ):
        assert want in texts, "题注文字被改动了：缺「%s」，实际：%s" % (want, texts)
    caps = [
        p
        for p in doc.paragraphs
        if p.text.strip().startswith(("表1.1", "表 １", "Figure 1-1", "Table 1-9"))
    ]
    assert len(caps) == 4, "题注未全部识别：%s" % [p.text[:20] for p in caps]
    assert all(p.alignment == WD_ALIGN_PARAGRAPH.CENTER for p in caps), "题注未居中"


def test_figure_caption_centered(tmp_path):
    """图注（Lua filter 标为 FigureCaption）居中，文字原样。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    md = "# 第一章 测试\n\n图 架构示意\n\n正文段落。\n"
    out = _build_md(tmp_path, md, {})
    doc = Document(out)
    cap = next(p for p in doc.paragraphs if p.text.strip().startswith("图 架构示意"))
    assert cap.alignment == WD_ALIGN_PARAGRAPH.CENTER, "图注未居中"


GRID_MD = (
    "# 第一章 测试\n\n"
    "表 1-1 多级表头\n\n"
    "+------------------+------------------+\n"
    "| 商务部分         | 技术部分         |\n"
    "+--------+---------+--------+---------+\n"
    "| 条款   | 响应    | 模块   | 说明   |\n"
    "+========+=========+========+=========+\n"
    "| 工期   | 完全响应| 接入层 | 设备   |\n"
    "+--------+---------+--------+---------+\n"
    "| 质保   | 优于要求| 平台层 | 治理   |\n"
    "+--------+---------+--------+---------+\n"
)


def _has_tbl_header(row):
    from docx.oxml.ns import qn

    trPr = row._tr.find(qn("w:trPr"))
    return trPr is not None and trPr.find(qn("w:tblHeader")) is not None


def test_grid_table_colspan_preserved(tmp_path):
    """grid table 的合并单元格要原样带过来。"""
    from docx import Document
    from docx.oxml.ns import qn

    out = _build_md(tmp_path, GRID_MD, {})
    row0 = Document(out).tables[0].rows[0]._tr.findall(qn("w:tc"))
    spans = []
    for tc in row0:
        tcPr = tc.find(qn("w:tcPr"))
        gs = tcPr.find(qn("w:gridSpan")) if tcPr is not None else None
        spans.append(gs.get(qn("w:val")) if gs is not None else None)
    assert spans == ["2", "2"], "合并单元格丢失：%s" % spans


def _shaded(cell):
    from docx.oxml.ns import qn

    tcPr = cell._tc.find(qn("w:tcPr"))
    sh = tcPr.find(qn("w:shd")) if tcPr is not None else None
    return sh is not None and sh.get(qn("w:fill")) == "EDEDED"


def test_pandoc_sets_tblheader_natively(tmp_path):
    """记录事实：tblHeader 是 pandoc 原生就给的（+===+ 以上全是表头行）。

    post.py 里的 set_repeat_header 只是幂等加固，不是这个功能的实现者。
    """
    from docx import Document

    out = _build_md(tmp_path, GRID_MD, {})
    rows = Document(out).tables[0].rows
    assert _has_tbl_header(rows[0]) and _has_tbl_header(rows[1])
    assert not _has_tbl_header(rows[2])


HEADERLESS_MD = (
    "# 第一章\n\n表 1-1 表单式表格\n\n"
    "+-----------+-----------+\n"
    "| 项目名称  | 某某项目  |\n"
    "+-----------+-----------+\n"
    "| 负责人    | 张三      |\n"
    "+-----------+-----------+\n"
)


def test_header_rows_zero_for_form_tables(tmp_path):
    """表单类表格：显式 header_rows=0 时首行不上灰底、也不跨页重复。

    （外部 docx 转 md 后 pandoc 会给首行加 tblHeader，表单的「项目名称」行
    看起来就成了列标题，视觉上不对。）
    """
    from docx import Document

    out = _build_md(tmp_path, HEADERLESS_MD, {"style": {"header_rows": 0}})
    tbl = Document(out).tables[0]
    assert not _shaded(tbl.rows[0].cells[0]), "header_rows=0 时首行不该有灰底"
    assert not _has_tbl_header(tbl.rows[0]), "header_rows=0 时不该跨页重复表头"


def test_auto_header_rows_from_pandoc(tmp_path):
    """默认自动识别：pandoc 标了几行表头，就给几行上灰底。"""
    from docx import Document

    out = _build_md(tmp_path, GRID_MD, {})
    rows = Document(out).tables[0].rows
    assert _shaded(rows[0].cells[0]) and _shaded(rows[1].cells[0]), "多级表头未全部灰底"
    assert not _shaded(rows[2].cells[0]), "数据行被误当成表头"


def test_auto_header_rows_per_table(tmp_path):
    """同一文档里不同表表头行数不同，必须各算各的（这是全局配置做不到的）。"""
    from docx import Document

    md = (
        "# 第一章\n\n"
        "| a | b |\n|:--|:--|\n| 1 | 2 |\n\n" + GRID_MD.split("# 第一章 测试\n\n", 1)[1]  # 单行表头
    )  # 两行表头
    out = _build_md(tmp_path, md, {})
    t0, t1 = Document(out).tables
    assert _shaded(t0.rows[0].cells[0]) and not _shaded(t0.rows[1].cells[0]), (
        "pipe table 只该第一行灰底"
    )
    assert _shaded(t1.rows[0].cells[0]) and _shaded(t1.rows[1].cells[0]), "grid table 该前两行灰底"


def test_header_rows_override(tmp_path):
    """显式给 header_rows 时以配置为准（覆盖自动识别）。"""
    from docx import Document

    out = _build_md(tmp_path, GRID_MD, {"style": {"header_rows": 1}})
    rows = Document(out).tables[0].rows
    assert _shaded(rows[0].cells[0])
    assert not _shaded(rows[1].cells[0]), "显式设为 1 时第二行不该有灰底"


def test_table_zebra_and_header_color(tmp_path):
    """表头可换深色底配白字，表体可加斑马纹（借用通用 docx 技能的表格规范）。

    默认（浅灰底 + 黑字 + 无斑马纹）走的是中文正式文档惯例，保持不动。
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    md = "# 第一章\n\n表 1-1 清单\n\n| 序号 | 名称 |\n|:--|:--|\n| 1 | a |\n| 2 | b |\n| 3 | c |\n"
    out = _build_md(
        tmp_path,
        md,
        {
            "style": {
                "table_shade": "4472C4",
                "table_header_color": "FFFFFF",
                "table_zebra": True,
                "table_zebra_fill": "F7F7F7",
            }
        },
    )
    tbl = Document(out).tables[0]

    def fill(ri):
        tcPr = tbl.rows[ri].cells[0]._tc.find(qn("w:tcPr"))
        shd = tcPr.find(qn("w:shd")) if tcPr is not None else None
        return shd.get(qn("w:fill")) if shd is not None else None

    assert fill(0) == "4472C4", "表头底纹未生效"
    run = tbl.rows[0].cells[0].paragraphs[0].runs[0]
    assert run.font.color.rgb == RGBColor(0xFF, 0xFF, 0xFF), "表头白字未生效"
    assert fill(1) is None, "首条数据行不该有斑马纹底"
    assert fill(2) == "F7F7F7", "第 2 条数据行应有斑马纹底"
    assert fill(3) is None, "斑马纹应隔行"


def test_table_style_defaults_unchanged(tmp_path):
    """不配 style 时保持中文惯例：浅灰表头、无斑马纹。"""
    from docx import Document
    from docx.oxml.ns import qn

    md = "# 第一章\n\n表 1-1 清单\n\n| 序号 | 名称 |\n|:--|:--|\n| 1 | a |\n| 2 | b |\n"
    tbl = Document(_build_md(tmp_path, md, {})).tables[0]

    def fill(ri):
        tcPr = tbl.rows[ri].cells[0]._tc.find(qn("w:tcPr"))
        shd = tcPr.find(qn("w:shd")) if tcPr is not None else None
        return shd.get(qn("w:fill")) if shd is not None else None

    assert fill(0) == "EDEDED", "默认表头应为浅灰"
    assert fill(1) is None, "默认不该有斑马纹"


def test_three_line_table(tmp_path):
    """三线表：内部无框线，顶底线加粗，表头行有下边框。"""
    from docx import Document
    from docx.oxml.ns import qn

    out = _build_md(tmp_path, GRID_MD, {"style": {"table_border": "three"}})
    tbl = Document(out).tables[0]
    borders = tbl._tbl.tblPr.find(qn("w:tblBorders"))

    def val(edge):
        el = borders.find(qn("w:" + edge))
        return el.get(qn("w:val")) if el is not None else None

    assert val("insideH") == "nil" and val("insideV") == "nil", "三线表内部不应有框线"
    assert val("top") == "single" and borders.find(qn("w:top")).get(qn("w:sz")) == "12"
    # 表头下线画在「最后一行表头」上；本表自动识别为 2 行表头，所以查 rows[1]
    last_head = tbl.rows[1].cells[0]._tc.find(qn("w:tcPr")).find(qn("w:tcBorders"))
    assert last_head is not None and last_head.find(qn("w:bottom")) is not None, "缺表头下线"


def test_caption_keeps_with_table(tmp_path):
    """表题不能与表格分家（keep_with_next）。"""
    from docx import Document

    out = _build_md(tmp_path, "# 第一章\n\n表 1-1 清单\n\n| a |\n|:--|\n| 1 |\n", {})
    cap = next(p for p in Document(out).paragraphs if p.text.strip().startswith("表 1-1"))
    assert cap.paragraph_format.keep_with_next, "表题未设 keep_with_next，可能与表格分页"


def test_page_number_default(tmp_path):
    """默认页码样式 — N —。"""
    from docx import Document

    out = _build_md(tmp_path, "# 第一章\n\n正文。\n", {})
    footer = Document(out).sections[-1].footer.paragraphs[0].text
    assert "—" in footer and "1" in footer, "默认页码异常：%r" % footer


def test_style_cfg_overrides(tmp_path):
    """config 的 style 段能改页码模板 / 题注颜色 / 字号。"""
    from docx import Document
    from docx.shared import RGBColor

    md = "# 第一章\n\n表 清单\n\n| a |\n|:--|\n| 1 |\n"
    cfg = {
        "style": {
            "page_number": "第 {n} 页",
            "toc_depth": "1-3",
            "caption_gray": "FF0000",
            "caption_size": 9.0,
        }
    }
    out = _build_md(tmp_path, md, cfg)
    doc = Document(out)

    footer = doc.sections[-1].footer.paragraphs[0].text
    assert "第" in footer and "页" in footer, "页码模板未生效：%r" % footer
    assert "—" not in footer, "旧页码样式残留：%r" % footer

    cap = next(p for p in doc.paragraphs if "清单" in p.text)
    assert cap.runs[0].font.color.rgb == RGBColor(0xFF, 0x00, 0x00), "题注颜色未生效"
    assert cap.runs[0].font.size.pt == 9.0, "题注字号未生效"


def test_style_cfg_numeric_and_spacing(tmp_path):
    """页眉/页码字号、单元格边距、边框、题注间距都要能被 style 段改。"""
    from docx import Document
    from docx.oxml.ns import qn

    md = "# 第一章\n\n表 清单\n\n| a | b |\n|:--|:--|\n| 1 | 2 |\n"
    out = _build_md(
        tmp_path,
        md,
        {
            "header": "页眉文字",
            "style": {
                "header_size": 12.0,
                "page_number_size": 14.0,
                "cell_margin_h": 200,
                "border_size": 12,
                "border_color": "000000",
                "caption_space_before": 12.0,
            },
        },
    )
    doc = Document(out)

    footer_run = doc.sections[-1].footer.paragraphs[0].runs[0]
    assert footer_run.font.size.pt == 14.0, "页码字号未生效"
    hdr = doc.sections[-1].header.paragraphs[0].runs[0]
    assert hdr.font.size.pt == 12.0, "页眉字号未生效"

    tbl = doc.tables[0]
    mar = tbl._tbl.tblPr.find(qn("w:tblCellMar"))
    assert mar.find(qn("w:left")).get(qn("w:w")) == "200", "单元格边距未生效"
    top = tbl._tbl.tblPr.find(qn("w:tblBorders")).find(qn("w:top"))
    assert top.get(qn("w:sz")) == "12" and top.get(qn("w:color")) == "000000"

    cap = next(p for p in doc.paragraphs if p.text.strip().startswith("表 清单"))
    assert cap.paragraph_format.space_before.pt == 12.0, "题注段前间距未生效"


def test_temp_dir_cleaned_on_failure(tmp_path):
    """渲染失败时也要清掉临时目录。

    以前只在 render() 末尾 rmtree，任何 sys.exit（配置有误、子进程报错）都会绕过它，
    临时目录就烂在 %TEMP% 里——实测一天攒了 12 个。
    """
    import glob
    import tempfile

    pat = os.path.join(tempfile.gettempdir(), "texere_*")
    before = set(glob.glob(pat))

    src = tmp_path / "src"
    src.mkdir()
    (src / "01.md").write_text("# 第一章\n\n正文。\n", encoding="utf-8")
    cfg = tmp_path / "cfg.json"
    cfg.write_text('{"content_fixes_file": "这个文件不存在.py"}', encoding="utf-8")

    r = subprocess.run(
        [
            sys.executable,
            os.path.join(KIT, "scripts", "render.py"),
            "--src",
            str(src),
            "--out",
            str(tmp_path / "o.docx"),
            "--config",
            str(cfg),
        ],
        capture_output=True,
        env=dict(os.environ, PYTHONIOENCODING="utf-8"),
    )
    assert r.returncode != 0, "这个配置应当失败"

    leaked = sorted(set(glob.glob(pat)) - before)
    assert not leaked, "失败后临时目录未清理：%s" % leaked


def test_render_finds_sibling_media(tmp_path):
    """图片在 src 的**兄弟**目录时也要能找到（真实项目就是这种布局：

    build/src/*.md 引用了 build/media_plan/*.jpg）。
    """
    pymupdf = pytest.importorskip("pymupdf")
    proj = tmp_path / "proj"
    (proj / "src").mkdir(parents=True)
    (proj / "media").mkdir()
    pymupdf.open().new_page().get_pixmap().save(str(proj / "media" / "img.png"))
    (proj / "src" / "01.md").write_text(
        "# 第一章\n\n![图 1-1 兄弟目录图片](media/img.png)\n\n正文。\n",
        encoding="utf-8",
    )

    out = str(tmp_path / "out.docx")
    r = subprocess.run(
        [
            sys.executable,
            os.path.join(KIT, "scripts", "render.py"),
            "--src",
            str(proj / "src"),
            "--out",
            out,
        ],
        capture_output=True,
        env=dict(os.environ, PYTHONIOENCODING="utf-8"),
    )
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")

    from docx import Document

    assert len(Document(out).inline_shapes) == 1, "兄弟目录里的图片没被找到"
    assert "images: 1/1 ok" in r.stdout.decode("utf-8", "replace"), "缺图自检未通过"


def test_images_survive_postprocess(tmp_path):
    """后处理绝不能把图片弄丢（曾被自动编号抹掉过 28 张）。"""
    pymupdf = pytest.importorskip("pymupdf")
    from docx import Document

    src = tmp_path / "src"
    src.mkdir(exist_ok=True)
    pymupdf.open().new_page().get_pixmap().save(str(src / "fig.png"))

    md = "# 第一章\n\n![图 架构示意](fig.png)\n\n正文引用。\n"
    out = _build_md(tmp_path, md, {})
    doc = Document(out)
    assert len(doc.inline_shapes) == 1, "图片被后处理弄丢了"


def test_caption_words_configurable(tmp_path):
    """题注关键字可配（post.py 与 lua filter 同步）。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    md = "# 第一章\n\n照片 架构示意\n\n正文段落。\n"
    words = {"table": ["清单"], "figure": ["照片"]}

    off = _build_md(tmp_path, md, {})
    para = next(p for p in Document(off).paragraphs if "照片" in p.text)
    assert para.alignment != WD_ALIGN_PARAGRAPH.CENTER, "默认关键字下不该把「照片」当图注"

    on = _build_md(tmp_path, md, {"caption_words": words})
    para = next(p for p in Document(on).paragraphs if "照片" in p.text)
    assert para.alignment == WD_ALIGN_PARAGRAPH.CENTER, "自定义关键字未生效（应被识别为图注并居中）"


def test_config_with_bom(tmp_path):
    """带 UTF-8 BOM 的 config.json 必须能读。

    Windows 记事本 / PowerShell 写出来的 json 默认带 BOM，真实用户高频踩到。
    """
    from docx import Document

    src = tmp_path / "src"
    src.mkdir(exist_ok=True)
    (src / "01.md").write_text("# 第一章\n\n正文。\n", encoding="utf-8")
    cfg = tmp_path / "cfg.json"
    cfg.write_bytes(b"\xef\xbb\xbf" + '{"header": "带 BOM 的页眉"}'.encode())

    body, out = str(tmp_path / "body.docx"), str(tmp_path / "out.docx")
    subprocess.run(_pandoc_cmd(str(src / "01.md"), body, str(src)), check=True, capture_output=True)
    r = subprocess.run(
        [sys.executable, os.path.join(KIT, "scripts", "post.py"), body, out, str(cfg)],
        capture_output=True,
        env=dict(os.environ, PYTHONIOENCODING="utf-8"),
    )
    assert r.returncode == 0, "带 BOM 的 config 读不了：" + r.stderr.decode("utf-8", "replace")
    assert "带 BOM 的页眉" in Document(out).sections[-1].header.paragraphs[0].text


def test_to_rgb_accepts_common_forms():
    import importlib.util

    spec = importlib.util.spec_from_file_location("post", os.path.join(KIT, "scripts", "post.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    assert mod.to_rgb("404040") == (0x40, 0x40, 0x40)
    assert mod.to_rgb("#FF0000") == (0xFF, 0, 0)
    assert mod.to_rgb("0x010203") == (1, 2, 3)
    assert mod.to_rgb([1, 2, 3]) == (1, 2, 3)
    with pytest.raises(ValueError):
        mod.to_rgb("abc")


def test_render_version_flag():
    r = subprocess.run(
        [sys.executable, os.path.join(KIT, "scripts", "render.py"), "--version"],
        capture_output=True,
    )
    assert r.returncode == 0
    assert b"texere" in r.stdout, r.stdout


def test_no_h1_is_processed(tmp_path):
    """表单/附件类文档没有一级标题也要能处理：跳过目录，且不产生空白首页。"""
    from docx import Document

    src = tmp_path / "src"
    src.mkdir()
    (src / "01.md").write_text(
        "项目名称：某某项目\n\n| a | b |\n|:--|:--|\n| 1 | 2 |\n", encoding="utf-8"
    )
    body, out = str(tmp_path / "body.docx"), str(tmp_path / "out.docx")
    subprocess.run(_pandoc_cmd(str(src / "01.md"), body, str(src)), check=True, capture_output=True)
    r = subprocess.run(
        [sys.executable, os.path.join(KIT, "scripts", "post.py"), body, out, ""],
        capture_output=True,
        env=dict(os.environ, PYTHONIOENCODING="utf-8"),
    )
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")
    doc = Document(out)
    assert len(doc.sections) == 1, "无封面无标题时不该分节（跳转符会变成空白首页）"
    assert "项目名称：某某项目" in [p.text.strip() for p in doc.paragraphs], "正文被丢了"
    assert not any(p.text.strip().startswith("目") and "录" in p.text for p in doc.paragraphs), (
        "没有标题就不该插目录"
    )
