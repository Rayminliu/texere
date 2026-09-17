# -*- coding: utf-8 -*-
"""通用 docx 后处理器：把 pandoc 产出的正文 docx 加工成正式中文文档。

用法: python post.py <body.docx> <out.docx> [config.json]

config.json 字段（均可省）:
  cover   : [[样式名, 文本], ...]  封面行，样式可用 CoverTop/CoverTitle/CoverSub/CoverInfo/CoverDate
  header  : 页眉文字（正文节）
  title / author / subject / comments : 文档属性
  toc_heading : 目录标题，默认 "目　　录"
行为: 去掉前置书名页与空段 -> 注入封面 -> 注入目录域 -> 分节(封面目录/正文各自页码)
      -> 正文节页眉页脚(居中页码) -> 表格 100% 宽/表头加粗灰底居中/单元格 10.5pt
      -> 表题与图注居中灰色去斜体 -> settings 加 updateFields。
所有手写 OOXML 均按 ECMA-376 子元素顺序插入，避免 Word 报"文档已损坏"。
"""
import copy
import json
import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SONG, LATIN = "宋体", "Times New Roman"
GRAY = (0x40, 0x40, 0x40)
# 可被 config.json 的 "style" 段覆盖（见 apply_style_cfg）。默认即中文正式文档惯例。
S = {
    "east_font": SONG,        # 中文字体
    "latin_font": LATIN,      # 西文字体
    "caption_gray": GRAY,     # 题注灰
    "caption_size": 10.5,     # 题注字号 pt
    "table_shade": "EDEDED",  # 表头底纹
    "table_size": 10.5,       # 表格字号 pt
    "toc_depth": "1-2",       # 目录收录层级
    "page_number": "— {n} —",  # 页码模板，{n} 处插入页码域
    "header_rows": None,      # 表头行数；None = 自动（读 pandoc 打的 w:tblHeader）
    "table_border": "full",   # 表格边框：full 全框线 / three 三线表 / none 无框线
}
# 兼容中文模板（reference_doc 来自中文 Word 时一级标题样式名为「标题 1」）
H1_STYLES = {"Heading 1", "标题 1"}
TITLE_STYLES = {"Title", "Subtitle", "Author", "Date", "标题", "副标题"}
# 题注编号：半角/全角数字 + 半角/全角分隔符
# 覆盖 表 1-1 / 表1.1 / 表１－１ / 图 2-3 / Table 1-1 / Figure 1-2
_D = r"[0-9０-９]"
_SEP = r"[.\-－—．]"
CAPTION_RE = re.compile(
    r"^(?:表|圖|图|表格|图片|Table|Figure|Fig\.?)\s*%s+(?:\s*%s\s*%s+)?" % (_D, _SEP, _D),
    re.IGNORECASE)
# 题注里的可选编号前缀，用于自动编号时替换（编号可写可不写）
CAP_NUM_RE = re.compile(
    (r"^(表|圖|图|表格|图片|Table|Figure|Fig\.?)"         # 1 题注关键字
     r"(\s*)"                                             # 2 分隔
     r"(%s+(?:\s*%s\s*%s+)?)?"                            # 3 可选旧编号
     r"(\s*)") % (_D, _SEP, _D),                          # 4 分隔
    re.IGNORECASE)
# 交叉引用标签：@tab:xxx / @fig:xxx（题注末尾声明，正文里引用）
CAP_LABEL_RE = re.compile(r"\s*@(tab|fig):([A-Za-z0-9_\-]+)\s*$")
REF_RE = re.compile(r"@(tab|fig):([A-Za-z0-9_\-]+)")
KIND_OF = {"表": "表", "圖": "图", "图": "图", "表格": "表", "图片": "图",
           "table": "表", "figure": "图", "fig": "图", "fig.": "图"}
# 由 filters/captions.lua 在 AST 层打上的语义样式（按 styleId 匹配，见 style_id 注释）
KIND_BY_STYLE_ID = {"TableCaption": "表", "FigureCaption": "图",
                    "ImageCaption": "图", "Caption": "图", "CaptionedFigure": "图"}
CAPTION_STYLE_IDS = set(KIND_BY_STYLE_ID)

PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
             "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
             "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
             "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
             "suppressOverlap", "jc", "textDirection", "textAlignment",
             "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
             "pPrChange"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
              "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
              "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange"]
TRPR_ORDER = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
              "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
               "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
               "tblCaption", "tblDescription"]
SETTINGS_ORDER = ["updateFields", "hdrShapeDefaults", "footnotePr", "endnotePr",
                  "compat", "rsids", "mathPr", "themeFontLang", "clrSchemeMapping",
                  "shapeDefaults", "decimalSymbol", "listSeparator"]


def insert_ordered(parent, child, order):
    tag = child.tag.split("}")[-1]
    idx = order.index(tag) if tag in order else len(order)
    for existing in parent:
        etag = existing.tag.split("}")[-1]
        if etag not in order:
            continue
        if order.index(etag) > idx:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = S["latin_font"]
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    rf.set(qn("w:ascii"), S["latin_font"])
    rf.set(qn("w:hAnsi"), S["latin_font"])
    rf.set(qn("w:eastAsia"), S["east_font"])


def add_field(paragraph, instr, placeholder="", size=10.5):
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    sp = OxmlElement("w:fldChar"); sp.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = placeholder
    en = OxmlElement("w:fldChar"); en.set(qn("w:fldCharType"), "end")
    for el in (b, it, sp, t, en):
        run._element.append(el)
    set_run_font(run, size=size)
    return run


def shade(cell, fill=None):
    fill = fill or S["table_shade"]
    tcpr = cell._tc.get_or_add_tcPr()
    sh = tcpr.find(qn("w:shd"))
    if sh is None:
        sh = insert_ordered(tcpr, OxmlElement("w:shd"), TCPR_ORDER)
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)


def set_cell_border(cell, edge, sz="6", color="000000"):
    tcPr = cell._tc.get_or_add_tcPr()
    tb = tcPr.find(qn("w:tcBorders"))
    if tb is None:
        tb = insert_ordered(tcPr, OxmlElement("w:tcBorders"), TCPR_ORDER)
    el = tb.find(qn("w:" + edge))
    if el is None:
        el = OxmlElement("w:" + edge)
        tb.append(el)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), sz)
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color)


def set_table_borders(tbl, mode, header_rows):
    """表格边框：full 全框线 / three 三线表 / none 无框线。

    三线表 = 顶线 + 表头下线 + 底线，内部无竖线无横线（论文/申报书常用）。
    """
    if mode not in ("full", "three", "none"):
        raise ValueError('table_border 只能是 "full" / "three" / "none"，收到：%r' % mode)
    tblPr = tbl._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = insert_ordered(tblPr, OxmlElement("w:tblBorders"), TBLPR_ORDER)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            borders.append(el)          # CT_TblBorders 有固定子元素顺序
        if mode == "none":
            el.set(qn("w:val"), "nil")
        elif mode == "three":
            if edge in ("top", "bottom"):
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "12")     # 1.5pt，三线表的顶底线要粗
                el.set(qn("w:space"), "0")
                el.set(qn("w:color"), "000000")
            else:
                el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "6")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "808080")
    if mode == "three":
        # 表头最后一行的下边框 = 三线表的中间那条线
        for cell in tbl.rows[min(header_rows, len(tbl.rows)) - 1].cells:
            set_cell_border(cell, "bottom")


def set_pgnum_start(sectPr, start=1):
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        anchor = sectPr.find(qn("w:cols")) or sectPr.find(qn("w:docGrid"))
        if anchor is not None:
            anchor.addprevious(pg)
        else:
            sectPr.append(pg)
    pg.set(qn("w:start"), str(start))


def set_repeat_header(row):
    """确保表头行跨页重复（w:tblHeader）。

    注意：pandoc 通常**已经**给表头行设了（grid table 里 `+===+` 以上的行都算表头），
    本函数只是幂等加固——若上游没设，表格跨页后第 2 页起就没有表头。
    """
    trPr = row._tr.find(qn("w:trPr"))
    if trPr is None:
        trPr = OxmlElement("w:trPr")
        row._tr.insert(0, trPr)
    el = trPr.find(qn("w:tblHeader"))
    if el is None:
        el = insert_ordered(trPr, OxmlElement("w:tblHeader"), TRPR_ORDER)
    el.set(qn("w:val"), "true")


def style_name(p):
    try:
        return p.style.name or ""
    except Exception:
        return ""


def style_id(p):
    """样式 styleId。比 name 可靠：内置「Table Caption」与我们的 TableCaption
    同名不同字，按 name 判会歧义，按 styleId 判唯一。"""
    try:
        return p.style.style_id or ""
    except Exception:
        return ""


def detect_header_rows(tbl):
    """数出开头连续带 w:tblHeader 的行数。

    pandoc 已经把表头行标好了（grid table 的 `+===+` 以上全部算表头），
    直接读它，比让用户填一个全局 header_rows 靠谱得多——同一文档里
    不同表的表头行数往往不一样。
    """
    n = 0
    for row in tbl.rows:
        trPr = row._tr.find(qn("w:trPr"))
        if trPr is None or trPr.find(qn("w:tblHeader")) is None:
            break
        n += 1
    return n


def to_rgb(v):
    """颜色解析：接受 "404040" / "#404040" / "0x404040" / [r,g,b]。"""
    if isinstance(v, (list, tuple)):
        return tuple(int(x) for x in v)
    s = str(v).strip().lstrip("#")
    if s.lower().startswith("0x"):
        s = s[2:]
    if len(s) != 6:
        raise ValueError("颜色应为 6 位十六进制（如 404040），收到：%r" % v)
    return tuple(int(s[i:i + 2], 16) for i in (0, 2, 4))


def apply_style_cfg(cfg):
    """把 config.json 的 "style" 段套到全局设置 S 上；没写的保持默认。"""
    st = cfg.get("style") or {}
    for k in ("east_font", "latin_font", "table_shade", "toc_depth", "page_number",
              "table_border"):
        if k in st:
            S[k] = st[k]
    for k in ("caption_size", "table_size"):
        if k in st:
            S[k] = float(st[k])
    if "header_rows" in st:
        S["header_rows"] = max(1, int(st["header_rows"]))
    if "caption_gray" in st:
        S["caption_gray"] = to_rgb(st["caption_gray"])
    return S


def find_h1(paras):
    """定位第一章标题；兼容中文模板的「标题 1」。找不到返回 None。"""
    for k, p in enumerate(paras):
        if style_name(p) in H1_STYLES:
            return k
    return None


def iter_blocks(doc):
    """按文档真实顺序遍历顶层段落与表格（doc.paragraphs 只给顶层段落，会漏掉表格）。"""
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def set_para_text(p, text):
    """整段文字写进首个 run 并清空其余 run。题注是纯文本，可安全合并。"""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def auto_number(doc):
    """表/图按章自动编号 + 交叉引用替换（config 里 "auto_number": true 启用）。

      表题  表 商务条款响应表 @tab:clause     ->  表 1-1 商务条款响应表
      图注  ![图 架构示意](a.png) @fig:arch   ->  图 1-2 架构示意
      引用  详见 @tab:clause                  ->  详见 表 1-1

    已手写的编号会被重排：插入/删除图表后无需人工对号。
    返回 (编号数, 引用数)。
    """
    chapter = 0
    counts = {"表": 0, "图": 0}
    labels = {}
    n_cap = 0

    for block in iter_blocks(doc):
        if isinstance(block, Table):
            continue                       # 单元格里的文字不参与编号
        sn = style_name(block)
        if sn in H1_STYLES:
            chapter += 1
            counts = {"表": 0, "图": 0}
            continue
        raw = block.text
        m_lab = CAP_LABEL_RE.search(raw)
        if m_lab:
            raw = raw[:m_lab.start()]
        txt = raw.strip()
        m = CAP_NUM_RE.match(txt)          # 有旧编号则匹配到，供下面剥离重排
        kind = KIND_BY_STYLE_ID.get(style_id(block))     # AST 层已标记：直接采信
        if kind is None:                   # 没走 filter 的文档才回落到文本判定
            if not m:
                continue
            kind = KIND_OF.get(m.group(1).lower(), "表")
        counts[kind] = counts.get(kind, 0) + 1
        number = "%s %d-%d" % (kind, chapter, counts[kind])
        body = txt[m.end():].strip() if m else txt
        set_para_text(block, "%s %s" % (number, body))
        if m_lab:
            labels[m_lab.group(2)] = number
        n_cap += 1

    n_ref = 0
    for block in iter_blocks(doc):
        # 表格要逐单元格取段落（python-docx 的 Table 没有 .paragraphs）
        paras = ([p for row in block.rows for cell in row.cells
                  for p in cell.paragraphs] if isinstance(block, Table) else [block])
        for p in paras:
            for r in p.runs:
                if "@" not in r.text:
                    continue
                new, k = REF_RE.subn(
                    lambda mm: labels.get(mm.group(2), mm.group(0)), r.text)
                if k:
                    r.text = new
                    n_ref += 1
    return n_cap, n_ref


def main(body_path, out_path, cfg_path):
    cfg = json.load(open(cfg_path, encoding="utf-8")) if cfg_path else {}
    apply_style_cfg(cfg)
    doc = Document(body_path)
    body = doc.element.body

    # 1. 定位第一章标题；删除前置书名页与空段
    h1_idx = find_h1(doc.paragraphs)
    if h1_idx is None:
        sys.exit("未找到任何一级标题（Markdown 的 # 标题）。post.py 靠它定位正文起点"
                 "以插入封面与目录；请检查源 md 是否含 # 标题，或 reference_doc 的"
                 "一级标题样式是否为 Heading 1 / 标题 1。")
    for p in list(doc.paragraphs[:h1_idx]):
        if style_name(p) in TITLE_STYLES:
            p._element.getparent().remove(p._element)
    while True:
        paras = doc.paragraphs
        h1_idx = find_h1(paras)
        if h1_idx == 0:
            break
        prev = paras[h1_idx - 1]
        if prev.text.strip() == "" and not prev._p.findall(".//" + qn("w:drawing")):
            prev._element.getparent().remove(prev._element)
        else:
            break
    first_h1 = doc.paragraphs[find_h1(doc.paragraphs)]
    first_h1.paragraph_format.page_break_before = False

    # 1.5 表/图自动编号（opt-in；必须在注入封面/目录之前，否则会把封面算进章序）
    n_auto = n_ref = 0
    if cfg.get("auto_number"):
        n_auto, n_ref = auto_number(doc)

    # 2. 封面 + 目录 + 分节段（先追加到末尾再整体前移）
    created = []

    def np(text="", style=None, align=None, page_break=False):
        p = doc.add_paragraph(text, style=style) if style else doc.add_paragraph(text)
        if align is not None:
            p.alignment = align
        if page_break:
            p.paragraph_format.page_break_before = True
        created.append(p)
        return p

    for style, text in cfg.get("cover", []):
        np(text, style)
    if cfg.get("cover"):
        np("", "CoverInfo")
    toc_head = np(cfg.get("toc_heading", "目　　录"), "TOC Heading",
                  align=WD_ALIGN_PARAGRAPH.CENTER, page_break=True)
    for r in toc_head.runs:
        set_run_font(r, size=16, bold=True, color=(0, 0, 0))
    toc_para = np("")
    toc_para.paragraph_format.first_line_indent = Pt(0)
    add_field(toc_para, 'TOC \\o "%s" \\h \\z \\u' % S["toc_depth"],
              "【目录将在打开文档时自动生成；若未显示请全选后按 F9】", size=12)
    sect_para = np("")
    sect_para.paragraph_format.first_line_indent = Pt(0)

    anchor = first_h1._p
    for p in created:
        anchor.addprevious(p._p)

    # 3. 分节与页码：封面+目录为第 1 节（无页眉页脚），正文为第 2 节
    body_sectPr = body.find(qn("w:sectPr"))
    sec1 = copy.deepcopy(body_sectPr)
    for tag in ("w:headerReference", "w:footerReference"):
        for el in sec1.findall(qn(tag)):
            sec1.remove(el)
    set_pgnum_start(sec1, 1)
    insert_ordered(sect_para._p.get_or_add_pPr(), sec1, PPR_ORDER)
    set_pgnum_start(body_sectPr, 1)

    sec_body = doc.sections[-1]
    sec_body.footer.is_linked_to_previous = False
    fp = sec_body.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Pt(0)
    left, _sep, right = S["page_number"].partition("{n}")
    if left:
        set_run_font(fp.add_run(left), size=9)
    add_field(fp, "PAGE", "1", size=9)
    if right:
        set_run_font(fp.add_run(right), size=9)

    header_text = cfg.get("header")
    if header_text:
        sec_body.header.is_linked_to_previous = False
        hp = sec_body.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.paragraph_format.first_line_indent = Pt(0)
        set_run_font(hp.add_run(header_text), size=9, color=(0x59, 0x59, 0x59))
        pbdr = OxmlElement("w:pBdr")
        btm = OxmlElement("w:bottom")
        btm.set(qn("w:val"), "single"); btm.set(qn("w:sz"), "4")
        btm.set(qn("w:space"), "1"); btm.set(qn("w:color"), "BFBFBF")
        pbdr.append(btm)
        insert_ordered(hp._p.get_or_add_pPr(), pbdr, PPR_ORDER)

    # 4. 表格排版
    for tbl in doc.tables:
        tblPr = tbl._tbl.tblPr
        w = tblPr.find(qn("w:tblW"))
        if w is None:
            w = insert_ordered(tblPr, OxmlElement("w:tblW"), TBLPR_ORDER)
        w.set(qn("w:type"), "pct"); w.set(qn("w:w"), "5000")
        layout = tblPr.find(qn("w:tblLayout"))
        if layout is None:
            layout = insert_ordered(tblPr, OxmlElement("w:tblLayout"), TBLPR_ORDER)
        layout.set(qn("w:type"), "autofit")
        mar = tblPr.find(qn("w:tblCellMar"))
        if mar is None:
            mar = insert_ordered(tblPr, OxmlElement("w:tblCellMar"), TBLPR_ORDER)
        for side, val in (("top", "40"), ("left", "80"), ("bottom", "40"), ("right", "80")):
            el = mar.find(qn("w:" + side))
            if el is None:
                el = OxmlElement("w:" + side); mar.append(el)
            el.set(qn("w:w"), val); el.set(qn("w:type"), "dxa")
        n_head = S["header_rows"] or detect_header_rows(tbl)
        header_rows = max(1, min(int(n_head), len(tbl.rows)))
        set_table_borders(tbl, S["table_border"], header_rows)
        for ri, row in enumerate(tbl.rows):
            is_head = ri < header_rows          # 多级表头时前 N 行都算表头
            if is_head:
                set_repeat_header(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if is_head:
                    shade(cell)
                for p in cell.paragraphs:
                    pf = p.paragraph_format
                    pf.first_line_indent = Pt(0)
                    pf.space_before = Pt(1); pf.space_after = Pt(1)
                    pf.line_spacing = 1.0
                    if is_head:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        set_run_font(r, size=S["table_size"],
                                     bold=True if is_head else None)

    # 5. 表题 / 图注 居中、灰色、去斜体
    n_cap = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if style_id(p) in CAPTION_STYLE_IDS or CAPTION_RE.match(txt):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.first_line_indent = Pt(0); pf.left_indent = Pt(0)
            pf.space_before = Pt(6); pf.space_after = Pt(4)
            # 表题必须与表格同页，否则会孤零零留在页尾
            pf.keep_with_next = True
            for r in p.runs:
                set_run_font(r, size=S["caption_size"], bold=False,
                             color=S["caption_gray"], italic=False)
            n_cap += 1

    # 6. 打开时自动刷域 + 文档属性
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
        insert_ordered(settings, uf, SETTINGS_ORDER)
    cp = doc.core_properties
    for k in ("title", "author", "subject", "comments"):
        if cfg.get(k):
            setattr(cp, k, cfg[k])

    doc.save(out_path)
    print("saved:", out_path, "| captions centered:", n_cap,
          "| tables:", len(doc.tables), "| sections:", len(doc.sections))
    if n_auto or n_ref:
        print("auto-number: %d captions, %d cross-refs" % (n_auto, n_ref))


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit("用法: python post.py <body.docx> <out.docx> [config.json]")
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)
