"""通用 docx 后处理器：把 pandoc 产出的正文 docx 加工成正式中文文档。

用法: python scripts/post.py <body.docx> <out.docx> [config.json]

config.json 字段（均可省）:
  cover   : [[样式名, 文本], ...]  封面行，样式可用 CoverTop/CoverTitle/CoverSub/CoverInfo/CoverDate
  header  : 页眉文字（正文节）
  title / author / subject / comments : 文档属性
  toc_heading : 目录标题，默认 "目　　录"
  toc   : false → 不插目录页（通知/公示类短文档）；标题样式照常保留，无封面时不分节
  mode  : "simple-report" → 短报告开箱模式（学生作业/短通知）：默认不插目录、
          所有 H1 不另起一页；缺省 "formal"。显式写的 toc / page_break_h1 优先于模式派生值
  page_break_h1 : false → 所有 H1 段前分页关闭（simple-report 自动派生）；
          缺省/formal = 第一章豁免（紧跟封面/目录），其余每章另起一页
  style   : 版式微调（字体/颜色/间距/表格边框等，见 README 的完整键表）
行为: 前置区处理（有封面时删 Title/Author 段，无封面时居中保留；目录插在标题块
      之后）-> 注入封面 -> 注入目录域 -> 分节(封面目录/正文各自页码)
      -> 正文节页眉页脚(居中页码) -> 表格 100% 宽/表头加粗灰底居中/单元格 10.5pt
      -> 表题与图注居中灰色去斜体 -> settings 加 updateFields。
契约: **只改版式，不改内容**——不触碰正文与题注的文字（图表编号由源文件手写）。
所有手写 OOXML 均按 ECMA-376 子元素顺序插入，避免 Word 报"文档已损坏"。
"""

import argparse
import copy
import json
import re
import sys

from _ooxml import (
    PPR_ORDER,
    SETTINGS_ORDER,
    TBLPR_ORDER,
    TCPR_ORDER,
    clear_repeat_header,
    insert_ordered,
    set_pgnum_start,
    set_repeat_header,
)
from _shared import (
    DEFAULT_CAPTION_WORDS,
    DEFAULT_EAST_FONT,
    DEFAULT_LATIN_FONT,
    H1_STYLE_ALIASES,
    H2_STYLE_ALIASES,
    H3_STYLE_ALIASES,
    PAGE_NUMBER_TEMPLATE,
    TITLE_STYLE_ALIASES,
    TOC_HEADING,
    TOC_PLACEHOLDER,
)
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GRAY = (0x40, 0x40, 0x40)
# 可被 config.json 的 "style" 段覆盖（见 apply_style_cfg）。默认即中文正式文档惯例；
# 语言相关默认值（字体对/页码模板/目录文案/题注词表）单一事实源在 _shared.py。
S = {
    "east_font": DEFAULT_EAST_FONT,  # 中文字体
    "latin_font": DEFAULT_LATIN_FONT,  # 西文字体
    "caption_gray": GRAY,  # 题注灰
    "caption_size": 10.5,  # 题注字号 pt
    "table_size": 10.5,  # 表格字号 pt
    "toc_depth": "1-2",  # 目录收录层级
    "page_number": PAGE_NUMBER_TEMPLATE,  # 页码模板，{n} 处插入页码域
    "header_rows": None,  # 表头行数；None = 自动（读 pandoc 打的 w:tblHeader）
    "table_border": "full",  # 表格边框：full 全框线 / three 三线表 / none 无框线
    # --- 页眉 / 页脚 ---
    "header_size": 9.0,  # 页眉字号 pt
    "header_gray": (0x59, 0x59, 0x59),
    "header_rule_color": "BFBFBF",  # 页眉下边框颜色
    "header_rule_size": 4,  # 页眉下边框粗细（1/8 pt）
    "page_number_size": 9.0,  # 页码字号 pt
    # --- 目录 ---
    "toc_title_size": 16.0,
    "toc_title_color": (0, 0, 0),
    "toc_placeholder": TOC_PLACEHOLDER,
    "toc_placeholder_size": 12.0,
    # --- 表格 ---
    "table_shade": "EDEDED",  # 表头底纹（配深色底时用 table_header_color 给白字）
    "table_header_color": None,  # 表头文字颜色；None = 不指定（继承黑色）
    "table_zebra": False,  # 斑马纹：表体隔行浅底
    "table_zebra_fill": "F7F7F7",
    "cell_margin_v": 40,  # 单元格上下边距 twips
    "cell_margin_h": 80,  # 单元格左右边距 twips
    "table_para_space": 1.0,  # 单元格内段前后 pt
    "border_size": 6,  # 全框线粗细（1/8 pt）
    "border_color": "808080",
    "three_line_size": 12,  # 三线表顶底线粗细（1/8 pt）
    # --- 题注 ---
    "caption_space_before": 6.0,
    "caption_space_after": 4.0,
    "caption_keep_with_next": True,  # 表题与表格同页；关掉可能省页数但会分家
    # --- 字号 / 页边距：模板基线的逐文档出口（make_ref 是全局模板，这里是单文档）
    # 只在 config.style 显式给出时才改写模板，未写的键原样保留甲方模板的自带版式
    "title_size": 26.0,  # 封面标题样式 Title（YAML title 的落点；作业类「标题三号」写 16）
    "body_size": 12.0,  # 正文 pt（Normal / Body Text / First Paragraph / Compact）
    "h1_size": 16.0,  # 一级标题 pt（Heading 1 / 标题 1）
    "h2_size": 14.0,
    "h3_size": 12.5,
    "margin_top": 2.54,  # cm（与 make_ref 模板基线一致）
    "margin_bottom": 2.54,
    "margin_left": 3.0,
    "margin_right": 2.6,
}
# 兼容中文模板（reference_doc 来自中文 Word 时一级标题样式名为「标题 1」）——别名表在 _shared
H1_STYLES = H1_STYLE_ALIASES
TITLE_STYLES = TITLE_STYLE_ALIASES
# 题注识别（文本兜底用；正常路径由 filters/captions.lua 在 AST 层打样式）
# 覆盖 表 1-1 / 表1.1 / 表１－１ / 图 2-3 / Table 1-1 / Figure 1-2；
# 关键字来自 _shared.DEFAULT_CAPTION_WORDS，与 captions.lua 默认词表逐词一致（守卫钉死）。
# 编号部分：全角/半角数字 + 多种分隔符（. - － — ．）
_D = r"[0-9０-９]"
_SEP = r"[.\-－—．]"


def _caption_re(words):
    """由词表构建题注匹配正则；长词优先，避免「表」抢「表格」的前缀。"""
    ws = sorted({w for lst in words.values() for w in lst if w}, key=len, reverse=True)
    alt = "|".join(re.escape(w) for w in ws)
    return re.compile(r"^(?:%s)\s*%s+(?:\s*%s\s*%s+)?" % (alt, _D, _SEP, _D), re.IGNORECASE)


def _table_prefix_re(table_words):
    r"""表题词「开头匹配」（不要求带编号）：keep_with_next 用，与历史硬编码
    ^\s*(表|表格|Table) 语义一致——不带编号的「表 商务条款响应表」也要粘住表格。"""
    ws = sorted({w for w in table_words if w}, key=len, reverse=True)
    alt = "|".join(re.escape(w) for w in ws)
    return re.compile(r"^\s*(?:%s)" % alt, re.IGNORECASE)


CAPTION_RE = _caption_re(DEFAULT_CAPTION_WORDS)
# 表题侧专用（keep_with_next 用）：默认/配置词表的 table 侧同源重建
TABLE_CAPTION_RE = _table_prefix_re(DEFAULT_CAPTION_WORDS["table"])
# 由 filters/captions.lua 在 AST 层打上的语义样式（按 styleId 匹配，见 style_id 注释）
CAPTION_STYLE_IDS = {
    "TableCaption",
    "FigureCaption",
    "ImageCaption",
    "Caption",
    "CaptionedFigure",
}


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = S["latin_font"]
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    rf.set(qn("w:ascii"), S["latin_font"])
    rf.set(qn("w:hAnsi"), S["latin_font"])
    rf.set(qn("w:eastAsia"), S["east_font"])


def add_field(paragraph, instr, placeholder="", size=10.5):
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    sp = OxmlElement("w:fldChar")
    sp.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = placeholder
    en = OxmlElement("w:fldChar")
    en.set(qn("w:fldCharType"), "end")
    for el in (b, it, sp, t, en):
        run._element.append(el)
    set_run_font(run, size=size)
    return run


def shade(cell, fill=None):
    fill = fill or S["table_shade"]
    tcpr = cell._tc.get_or_add_tcPr()
    sh = tcpr.find(qn("w:shd"))
    if sh is None:
        sh = insert_ordered(tcpr, OxmlElement("w:shd"), TCPR_ORDER)
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)


def set_cell_border(cell, edge, sz=None, color=None):
    sz = str(sz if sz is not None else S["border_size"])
    color = color or S["border_color"]
    tcPr = cell._tc.get_or_add_tcPr()
    tb = tcPr.find(qn("w:tcBorders"))
    if tb is None:
        tb = insert_ordered(tcPr, OxmlElement("w:tcBorders"), TCPR_ORDER)
    el = tb.find(qn("w:" + edge))
    if el is None:
        el = OxmlElement("w:" + edge)
        tb.append(el)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), sz)
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color)


def set_table_borders(tbl, mode, header_rows):
    """表格边框：full 全框线 / three 三线表 / none 无框线。

    三线表 = 顶线 + 表头下线 + 底线，内部无竖线无横线（论文/申报书常用）。
    """
    if mode not in ("full", "three", "none"):
        raise ValueError('table_border 只能是 "full" / "three" / "none"，收到：%r' % mode)
    tblPr = tbl._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = insert_ordered(tblPr, OxmlElement("w:tblBorders"), TBLPR_ORDER)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            borders.append(el)  # CT_TblBorders 有固定子元素顺序
        if mode == "none":
            el.set(qn("w:val"), "nil")
        elif mode == "three":
            if edge in ("top", "bottom"):
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), str(S["three_line_size"]))  # 顶底线比内部线粗
                el.set(qn("w:space"), "0")
                el.set(qn("w:color"), "000000")
            else:
                el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(S["border_size"]))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), S["border_color"])
    if mode == "three" and header_rows >= 1:
        # 表头最后一行的下边框 = 三线表的中间那条线（没有表头行就不画）
        for cell in tbl.rows[min(header_rows, len(tbl.rows)) - 1].cells:
            set_cell_border(cell, "bottom")


def style_name(p):
    try:
        return p.style.name or ""
    except Exception:
        return ""


def style_id(p):
    """样式 styleId。比 name 可靠：内置「Table Caption」与我们的 TableCaption
    同名不同字，按 name 判会歧义，按 styleId 判唯一。"""
    try:
        return p.style.style_id or ""
    except Exception:
        return ""


def build_caption_matchers(words):
    """按 config 的 caption_words 重建题注关键字（默认词表见 _shared）。

    words 形如 {"table": ["表", "表格"], "figure": ["图", "图片"]}。
    同时供 filters/captions.lua 使用（由 render.py 以 -M 传入），两边保持一致。
    表题侧 TABLE_CAPTION_RE 单独重建：keep_with_next 只认表题词，不跟着 figure 词漂。
    """
    global CAPTION_RE, TABLE_CAPTION_RE
    ws = (words.get("table") or []) + (words.get("figure") or [])
    ws = [w for w in ws if w]
    if not ws:
        return
    CAPTION_RE = _caption_re({"all": ws})
    if words.get("table"):
        TABLE_CAPTION_RE = _table_prefix_re(words["table"])


def uses_caption_styles(doc):
    """文档里是否已有 AST 层打的题注样式。

    有的话就**只信样式**，不再用文本正则兜底——否则「**表层（边缘轻算力）：**基于…」
    这类正文会被当成表题，插入编号并毁掉加粗（真实项目实测发生过）。
    """
    for p in doc.paragraphs:
        if style_id(p) in CAPTION_STYLE_IDS:
            return True
    return False


def detect_header_rows(tbl):
    """数出开头连续带 w:tblHeader 的行数。

    pandoc 已经把表头行标好了（grid table 的 `+===+` 以上全部算表头），
    直接读它，比让用户填一个全局 header_rows 靠谱得多——同一文档里
    不同表的表头行数往往不一样。
    """
    n = 0
    for row in tbl.rows:
        trPr = row._tr.find(qn("w:trPr"))
        if trPr is None or trPr.find(qn("w:tblHeader")) is None:
            break
        n += 1
    return n


def to_rgb(v):
    """颜色解析：接受 "404040" / "#404040" / "0x404040" / [r,g,b]。"""
    if isinstance(v, (list, tuple)):
        return tuple(int(x) for x in v)
    s = str(v).strip().lstrip("#")
    if s.lower().startswith("0x"):
        s = s[2:]
    if len(s) != 6:
        raise ValueError("颜色应为 6 位十六进制（如 404040），收到：%r" % v)
    return tuple(int(s[i : i + 2], 16) for i in (0, 2, 4))


# style 段各键的类型（决定怎么解析），见 apply_style_cfg
_STYLE_STR = (
    "east_font",
    "latin_font",
    "toc_depth",
    "page_number",
    "table_border",
    "toc_placeholder",
    "table_shade",
    "table_zebra_fill",
    "header_rule_color",
    "border_color",
)
_STYLE_FLOAT = (
    "caption_size",
    "table_size",
    "header_size",
    "page_number_size",
    "toc_title_size",
    "toc_placeholder_size",
    "table_para_space",
    "caption_space_before",
    "caption_space_after",
    "body_size",
    "h1_size",
    "h2_size",
    "h3_size",
    "title_size",
    "margin_top",
    "margin_bottom",
    "margin_left",
    "margin_right",
)
_STYLE_INT = (
    "cell_margin_v",
    "cell_margin_h",
    "border_size",
    "three_line_size",
    "header_rule_size",
)
_STYLE_COLOR = ("caption_gray", "header_gray", "toc_title_color", "table_header_color")

# style 段全部合法键：四张类型清单 + apply_style_cfg 显式特判的键 + caption_words
# （语义键走兼容路径，不是样式但 style 里接受）。与 config.schema.json 的 style
# 键集合必须一致——test_postprocess 双向钉死；漏进白名单的键会被未知键 [warn] 误伤。
KNOWN_STYLE_KEYS = frozenset(_STYLE_STR + _STYLE_FLOAT + _STYLE_INT + _STYLE_COLOR) | {
    "header_rows",
    "caption_keep_with_next",
    "table_zebra",
    "caption_words",
}


def apply_style_cfg(cfg):
    """把 config.json 的 "style" 段套到全局设置 S 上；没写的保持默认。

    未知键发 [warn]：拼错的键（body_ize / h1_ize）此前会**静默失效**——顶层
    段有同样的 [warn] 机制（render._load_config），style 子键不能成为盲区。
    """
    st = cfg.get("style") or {}
    unknown = sorted(set(st) - KNOWN_STYLE_KEYS)
    if unknown:
        print(
            "[warn] style 段键 %s 不被识别（多半是拼写有误，如 body_size 写成 body_ize）——"
            "该键会被忽略，键表见 docs/CONFIG.zh-CN.md" % ", ".join(repr(k) for k in unknown)
        )
    for k in _STYLE_STR:
        if k in st:
            S[k] = st[k]
    for k in _STYLE_FLOAT:
        if k in st:
            S[k] = float(st[k])
    for k in _STYLE_INT:
        if k in st:
            S[k] = int(st[k])
    for k in _STYLE_COLOR:
        if st.get(k) is not None:  # 允许显式写 null = 恢复默认（不指定）
            S[k] = to_rgb(st[k])
    if "header_rows" in st:
        # 允许显式写 0 = 这张表没有表头（表单/附件类表格），不要灰底、不要重复表头
        S["header_rows"] = max(0, int(st["header_rows"]))
    if "caption_keep_with_next" in st:
        S["caption_keep_with_next"] = bool(st["caption_keep_with_next"])
    if "table_zebra" in st:
        S["table_zebra"] = bool(st["table_zebra"])
    # caption_words 是顶层键（它不是"样式"，是语义），也兼容写在 style 里
    cw = cfg.get("caption_words") or st.get("caption_words")
    if cw:
        build_caption_matchers(cw)
    return S


def _normalize_mode(cfg):
    """mode 归一：simple-report 派生 toc / page_break_h1 的默认值；非法值立刻报错。

    显式写的键优先于模式派生值——用户在 simple-report 里写 "toc": true 就该有目录。
    """
    mode = cfg.get("mode", "formal")
    if mode not in ("formal", "simple-report"):
        sys.exit('mode 只支持 "formal" / "simple-report"，收到：%r' % mode)
    if "toc" not in cfg:
        cfg["toc"] = mode != "simple-report"
    if "page_break_h1" not in cfg:
        cfg["page_break_h1"] = mode != "simple-report"
    return mode


_BODY_STYLE_NAMES = ("Normal", "Body Text", "First Paragraph", "Compact")
_HEAD_SIZE_KEYS = {
    "h1_size": H1_STYLE_ALIASES,
    "h2_size": H2_STYLE_ALIASES,
    "h3_size": H3_STYLE_ALIASES,
    # 只 Title 本样式，不含 TITLE_STYLE_ALIASES 里的 Subtitle/Author/Date——
    # 作者行不该跟着标题字号走；那些要调走 custom-style 或模板
    "title_size": {"Title", "标题"},
}
_MARGIN_KEYS = ("margin_top", "margin_bottom", "margin_left", "margin_right")


def _apply_style_overrides(doc, style_cfg):
    """字号/页边距落到 styles.xml 与 sectPr：模板给基线，这里给逐文档出口。

    用户不该为了改个字号去重建 ref.docx（R4 #5）。只动显式给出的键：没写的键
    保留模板原值——甲方模板（reference_doc）自带的页边距/字号不被无意识冲掉。
    """
    if not style_cfg:
        return
    if "body_size" in style_cfg:
        size = float(style_cfg["body_size"])
        for name in _BODY_STYLE_NAMES:
            try:
                doc.styles[name].font.size = Pt(size)
            except KeyError:
                continue
        # 首行缩进跟着字号走：模板按 12pt × 2 字符算死在样式里，字号变了要重算
        for name in ("Body Text", "First Paragraph"):
            try:
                doc.styles[name].paragraph_format.first_line_indent = Pt(size * 2)
            except KeyError:
                continue
    for key, names in _HEAD_SIZE_KEYS.items():
        if key not in style_cfg:
            continue
        for name in names:
            try:
                doc.styles[name].font.size = Pt(float(style_cfg[key]))
            except KeyError:
                continue
    for key in _MARGIN_KEYS:
        if key not in style_cfg:
            continue
        side = key.split("_")[1]  # margin_left -> left
        for sec in doc.sections:
            setattr(sec, side + "_margin", Cm(float(style_cfg[key])))


def find_h1(paras):
    """定位第一章标题；兼容中文模板的「标题 1」。找不到返回 None。"""
    for k, p in enumerate(paras):
        if style_name(p) in H1_STYLES:
            return k
    return None


def _strip_front_matter(doc, cfg):
    """前置区处理：有封面时删 Title/Author 段，无封面时居中保留；返回 H1 下标。

    - 有 cover：pandoc 从 YAML 生成的 Title/Author/Date 段与封面行重复，删掉。
    - 无 cover：标题块就是文档门面，居中保留，目录/封面会插在它之后——此前
      一律删除，短文档用户只能靠 custom-style 绕行（R4 #2 实测）。
    - 表单/附件类文档没有一级标题也要能处理：跳过目录（没有标题可索引），
      封面照样插到最前面，表格与题注排版照做。
    """
    h1_idx = find_h1(doc.paragraphs)
    if h1_idx is None:
        print("[warn] 未找到一级标题（表单/附件类文档常见）：")
        print("       将跳过目录注入，只做封面与表格/题注排版。")
    else:
        for p in list(doc.paragraphs[:h1_idx]):
            if style_name(p) in TITLE_STYLES:
                if cfg.get("cover"):
                    p._element.getparent().remove(p._element)
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        while True:
            paras = doc.paragraphs
            h1_idx = find_h1(paras)
            if h1_idx == 0:
                break
            prev = paras[h1_idx - 1]
            if prev.text.strip() == "" and not prev._p.findall(".//" + qn("w:drawing")):
                prev._element.getparent().remove(prev._element)
            else:
                break
        h1_idx = find_h1(doc.paragraphs)

        # H1 段前分页（R4 #3）：False = 全部关闭（simple-report 模式派生 / 显式配置）；
        # 否则保持模板的「每章另起一页」，但第一章仍豁免——它紧跟封面/目录，
        # 另起一页只会剩下大半页空白。
        if cfg.get("page_break_h1") is False:
            for p in doc.paragraphs:
                if style_name(p) in H1_STYLES:
                    p.paragraph_format.page_break_before = False
        else:
            doc.paragraphs[h1_idx].paragraph_format.page_break_before = False

        # 源文件在第一个 # 之前还写了正文段落时（标题块不算）：toc 开着它们会
        # 排在目录之后。只提示、不删除——本工具的契约是不改内容。
        leftover = [
            p.text.strip()
            for p in doc.paragraphs[:h1_idx]
            if p.text.strip() and style_name(p) not in TITLE_STYLES
        ]
        if leftover:
            if cfg.get("toc", True) is False:
                print(
                    "[info] toc 已关闭：第一个一级标题之前的 %d 段内容保留在文档开头"
                    % len(leftover)
                )
            else:
                print("[warn] 第一个一级标题之前还有 %d 段内容，会排在目录之后：" % len(leftover))
                for t in leftover[:3]:
                    print("       " + t[:60])
                print("       建议：从源文件删掉，或写进 config 的 cover 由封面承载")

    return h1_idx


def _inject_front_pages(doc, body, cfg, h1_idx):
    """按 config 生成封面段与目录域，插到正文最前面。

    返回 (分节占位段, 是否需要第 1 节)——真正的分节由 _setup_sections 落地。
    """
    created = []

    def np(text="", style=None, align=None, page_break=False):
        p = doc.add_paragraph(text, style=style) if style else doc.add_paragraph(text)
        if align is not None:
            p.alignment = align
        if page_break:
            p.paragraph_format.page_break_before = True
        created.append(p)
        return p

    for style, text in cfg.get("cover", []):
        np(text, style)
    if cfg.get("cover"):
        np("", "CoverInfo")
    # "toc": false —— 短文档（通知/公示）不要目录页但要标题样式。
    # 旧写法只能靠「不写 #」绕过，代价是全文变普通段落、标题样式手工后补。
    toc_enabled = cfg.get("toc", True) is not False
    if h1_idx is not None and toc_enabled:  # 没有一级标题就不插目录（无处可索引）
        toc_head = np(
            cfg.get("toc_heading", TOC_HEADING),
            "TOC Heading",
            align=WD_ALIGN_PARAGRAPH.CENTER,
            page_break=True,
        )
        for r in toc_head.runs:
            set_run_font(r, size=S["toc_title_size"], bold=True, color=S["toc_title_color"])
        toc_para = np("")
        toc_para.paragraph_format.first_line_indent = Pt(0)
        add_field(
            toc_para,
            'TOC \\o "%s" \\h \\z \\u' % S["toc_depth"],
            S["toc_placeholder"],
            size=S["toc_placeholder_size"],
        )
    # 只有真的往第 1 节里放了东西（封面或目录）才分节；否则那个空的分节段
    # 会变成一张完全空白的首页（表单类文档实测踩过）。
    create_sec1 = bool(cfg.get("cover")) or (h1_idx is not None and toc_enabled)
    sect_para = None
    if create_sec1:
        sect_para = np("")
        sect_para.paragraph_format.first_line_indent = Pt(0)

    # 封面必须是文档第一页：插到 body 最前面，而不是「第一个标题之前」。
    # 否则源文件在第一个 # 之前写的内容会排到封面之前，单独占一页（实测踩过）。
    # 无封面且标题块被保留时（_strip_front_matter 只居中不删），插入点退到标题块
    # 之后——标题是文档门面，必须排在目录前面。
    anchor = next((c for c in body.iterchildren() if c.tag in (qn("w:p"), qn("w:tbl"))), None)
    if not cfg.get("cover"):
        last_title_el = None
        for p in doc.paragraphs:
            if style_name(p) in TITLE_STYLES:
                last_title_el = p._p
            else:
                break
        if last_title_el is not None:
            nxt = last_title_el.getnext()
            while nxt is not None and nxt.tag not in (qn("w:p"), qn("w:tbl")):
                nxt = nxt.getnext()
            if nxt is not None:
                anchor = nxt
    if anchor is None:
        # 空正文（无段无表）：回落到 sectPr 前插入
        sect_el = body.find(qn("w:sectPr"))
        anchor = sect_el if sect_el is not None else body.makeelement(qn("w:p"), {})
        if sect_el is None:
            body.append(anchor)
    for p in created:
        anchor.addprevious(p._p)

    return sect_para, create_sec1


def _setup_sections(doc, body, cfg, sect_para, create_sec1):
    """分节与页码：封面+目录为第 1 节（无页眉页脚），正文为第 2 节。"""
    body_sectPr = body.find(qn("w:sectPr"))
    if create_sec1:
        sec1 = copy.deepcopy(body_sectPr)
        for tag in ("w:headerReference", "w:footerReference"):
            for el in sec1.findall(qn(tag)):
                sec1.remove(el)
        set_pgnum_start(sec1, 1)
        insert_ordered(sect_para._p.get_or_add_pPr(), sec1, PPR_ORDER)
    set_pgnum_start(body_sectPr, 1)

    sec_body = doc.sections[-1]
    if S["page_number"] is None:
        # "page_number": null —— 甲方模板的页脚常带公司名 / 文档编号，别把它冲掉。
        # 这里一个字都不写，让 reference_doc 自带的页脚原样留下。
        print("      page_number: null -> 保留模板自带页脚")
    else:
        sec_body.footer.is_linked_to_previous = False
        fp = sec_body.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.first_line_indent = Pt(0)
        left, _sep, right = S["page_number"].partition("{n}")
        if left:
            set_run_font(fp.add_run(left), size=S["page_number_size"])
        add_field(fp, "PAGE", "1", size=S["page_number_size"])
        if right:
            set_run_font(fp.add_run(right), size=S["page_number_size"])

    header_text = cfg.get("header")
    if header_text:
        sec_body.header.is_linked_to_previous = False
        hp = sec_body.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.paragraph_format.first_line_indent = Pt(0)
        set_run_font(hp.add_run(header_text), size=S["header_size"], color=S["header_gray"])
        pbdr = OxmlElement("w:pBdr")
        btm = OxmlElement("w:bottom")
        btm.set(qn("w:val"), "single")
        btm.set(qn("w:sz"), str(S["header_rule_size"]))
        btm.set(qn("w:space"), "1")
        btm.set(qn("w:color"), S["header_rule_color"])
        pbdr.append(btm)
        insert_ordered(hp._p.get_or_add_pPr(), pbdr, PPR_ORDER)


def _format_tables(doc):
    """表格排版：宽度 100%、autofit、单元格边距、表头/斑马纹、字号。只改版式。"""
    for tbl in doc.tables:
        tblPr = tbl._tbl.tblPr
        w = tblPr.find(qn("w:tblW"))
        if w is None:
            w = insert_ordered(tblPr, OxmlElement("w:tblW"), TBLPR_ORDER)
        w.set(qn("w:type"), "pct")
        w.set(qn("w:w"), "5000")
        layout = tblPr.find(qn("w:tblLayout"))
        if layout is None:
            layout = insert_ordered(tblPr, OxmlElement("w:tblLayout"), TBLPR_ORDER)
        layout.set(qn("w:type"), "autofit")
        mar = tblPr.find(qn("w:tblCellMar"))
        if mar is None:
            mar = insert_ordered(tblPr, OxmlElement("w:tblCellMar"), TBLPR_ORDER)
        for side, val in (
            ("top", S["cell_margin_v"]),
            ("left", S["cell_margin_h"]),
            ("bottom", S["cell_margin_v"]),
            ("right", S["cell_margin_h"]),
        ):
            el = mar.find(qn("w:" + side))
            if el is None:
                el = OxmlElement("w:" + side)
                mar.append(el)
            el.set(qn("w:w"), str(val))
            el.set(qn("w:type"), "dxa")
        # 表头行数：显式配置优先，否则读 pandoc 打的 tblHeader；
        # 都为 0（表单类表格没有表头）时就是 0——绝不能强设成 1，否则首行会被误上灰底
        n_head = S["header_rows"] if S["header_rows"] is not None else detect_header_rows(tbl)
        header_rows = max(0, min(int(n_head), len(tbl.rows)))
        if S["header_rows"] == 0:  # 显式声明无表头：连「跨页重复」也去掉
            for row in tbl.rows:
                clear_repeat_header(row)
        set_table_borders(tbl, S["table_border"], header_rows)
        for ri, row in enumerate(tbl.rows):
            is_head = ri < header_rows  # 多级表头时前 N 行都算表头
            # 斑马纹：表体从第一条数据行起隔行浅底（首条数据行保持白底）
            zebra = bool(S["table_zebra"]) and not is_head and (ri - header_rows) % 2 == 1
            if is_head:
                set_repeat_header(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if is_head:
                    shade(cell)
                elif zebra:
                    shade(cell, S["table_zebra_fill"])
                for p in cell.paragraphs:
                    pf = p.paragraph_format
                    pf.first_line_indent = Pt(0)
                    pf.space_before = Pt(S["table_para_space"])
                    pf.space_after = Pt(S["table_para_space"])
                    pf.line_spacing = 1.0
                    if is_head:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        set_run_font(
                            r,
                            size=S["table_size"],
                            bold=True if is_head else None,
                            color=S["table_header_color"] if is_head else None,
                        )


def _format_captions(doc):
    """表题/图注居中、灰色、去斜体；返回处理条数（main 的人读输出要报）。"""
    n_cap = 0
    strict = uses_caption_styles(doc)  # 同上：有样式就不靠正则猜
    for p in doc.paragraphs:
        txt = p.text.strip()
        if style_id(p) in CAPTION_STYLE_IDS or (not strict and CAPTION_RE.match(txt)):
            sid = style_id(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.first_line_indent = Pt(0)
            pf.left_indent = Pt(0)
            pf.space_before = Pt(S["caption_space_before"])
            pf.space_after = Pt(S["caption_space_after"])
            # 「与下一段同页」只对**位于表格之前**的表题、和**载有图片**的段落有意义。
            # 图注本身在图片之后，粘住它会把后面的内容整块推走：实测 47 图文档多出 4 页。
            keep = bool(S["caption_keep_with_next"]) and (
                sid in ("TableCaption", "CaptionedFigure") or bool(TABLE_CAPTION_RE.match(txt))
            )
            pf.keep_with_next = keep
            for r in p.runs:
                set_run_font(
                    r,
                    size=S["caption_size"],
                    bold=False,
                    color=S["caption_gray"],
                    italic=False,
                )
            n_cap += 1

    return n_cap


def _set_update_fields_and_props(doc, cfg):
    """打开时自动刷域 + 写文档属性。"""
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        insert_ordered(settings, uf, SETTINGS_ORDER)
    cp = doc.core_properties
    for k in ("title", "author", "subject", "comments"):
        if cfg.get(k):
            setattr(cp, k, cfg[k])


def main(body_path, out_path, cfg_path):
    """后处理编排：清前置 → 插封面/目录 → 分节 → 表格 → 题注 → 属性。

    每一步都是只改版式、不碰文字的小函数（契约见模块 docstring）。
    """
    cfg = {}
    if cfg_path:
        with open(cfg_path, encoding="utf-8-sig") as f:
            cfg = json.load(f)
    apply_style_cfg(cfg)
    _normalize_mode(cfg)
    doc = Document(body_path)
    body = doc.element.body
    _apply_style_overrides(doc, cfg.get("style") or {})

    h1_idx = _strip_front_matter(doc, cfg)
    sect_para, create_sec1 = _inject_front_pages(doc, body, cfg, h1_idx)
    _setup_sections(doc, body, cfg, sect_para, create_sec1)
    _format_tables(doc)
    n_cap = _format_captions(doc)
    _set_update_fields_and_props(doc, cfg)

    doc.save(out_path)
    print(
        "saved:",
        out_path,
        "| captions centered:",
        n_cap,
        "| tables:",
        len(doc.tables),
        "| sections:",
        len(doc.sections),
    )


def _build_parser():
    p = argparse.ArgumentParser(
        prog="post.py",
        description="对 pandoc 产出的正文 docx 做后处理：封面/目录注入、分节、表格与题注排版。",
        epilog="契约：只改版式，不改内容（不触碰正文与题注的文字）",
    )
    p.add_argument("body_path", help="pandoc 产出的正文 docx")
    p.add_argument("out_path", help="后处理结果写出的路径")
    p.add_argument("config", nargs="?", help="样式配置 json（可省，见模块 docstring 的字段表）")
    return p


if __name__ == "__main__":
    a = _build_parser().parse_args()
    main(a.body_path, a.out_path, a.config)
