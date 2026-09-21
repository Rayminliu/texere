"""编辑已有 docx：定点修改，其余部分原样保留。

与 render.py 的生成链路互补，契约正好相反：

  生成链路  md -> pandoc -> 全新文档   只改版式，不改内容
  编辑链路  docx -> 定点改 -> 存回     只改你指定的地方，其余字节不动
            （不注入封面 / 目录 / 页码 / 不重排样式）

用法:
  python scripts/edit.py 标书.docx --list
  python scripts/edit.py 标书.docx --replace "旧=新" [--replace "旧2=新2"]
  python scripts/edit.py 标书.docx --after  "锚点文字" --text "新增段落"
  python scripts/edit.py 标书.docx --before "锚点文字" --text "新增段落"
  python scripts/edit.py 标书.docx --delete "待删除段落所含文字"
  python scripts/edit.py 标书.docx --cell 0 2 1 "1,060,000"
  python scripts/edit.py 标书.docx --add-row 0 "接入层" "设备" "320,000"
  python scripts/edit.py 标书.docx --add-rows 0 3 [--template-row 2]   # 复制行格式批量加空行
  python scripts/edit.py 标书.docx --fill data.json                    # 批量填表（JSON/CSV）
  python scripts/edit.py 标书.docx --del-row 0 2
  python scripts/edit.py 标书.docx --header "新版页眉" [--section body|all|N]
  python scripts/edit.py 标书.docx --footer "第 X 页"  [--section body|all|N]
  python scripts/edit.py 标书.docx --replace "A=B" --verify

默认写回原文件并先备份成 <name>.bak.docx；用 --out 另存，用 --no-backup 关掉备份。
--verify 会调 finalize.py 让 Word 真正打开一次（打不开 = 改坏了结构），复用生成链路的验收。

已知边界（不做）：批注 / 修订 / 水印 / 内容控件 / 含宏的 .docm。
"""

import argparse
import copy
import os
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

KIT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Windows 控制台默认 GBK，遇到无法编码的字符会抛 UnicodeEncodeError 让整条链崩掉。
# 保持控制台原编码不变，只把无法编码的字符降级为 ?（与 render.py 同一取舍）。
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(errors="replace")

XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


# ---------------------------------------------------------------- 基础工具


def _t_nodes(p):
    """段落内全部 w:t 节点（含超链接里的文字），按文档顺序。"""
    return p._p.findall(".//" + qn("w:t"))


def _set_t(node, text):
    """写 w:t 文本；首尾有空白时必须加 xml:space="preserve"，否则 Word 会吞掉空格。"""
    node.text = text
    if text[:1].isspace() or text[-1:].isspace():
        node.set(XML_SPACE, "preserve")
    else:
        node.attrib.pop(XML_SPACE, None)


def _has_graphic(p):
    """段落里含图片 / 图形对象时返回 True。契约要求不破坏图，遇到就跳过。"""
    for tag in ("w:drawing", "w:pict", "w:object"):
        if p._p.find(".//" + qn(tag)) is not None:
            return True
    return False


def _node_new_text(seg, off, chosen):
    """按原有 run 切分，算出这个 run 替换后应该是什么文本。

    seg     : 该 run 原本的文本
    off     : 该 run 在整段文本中的起始偏移
    chosen  : 已去重叠的替换区间 [(start, end, new), ...]

    关键点：替换内容只写进「区间起点所在的那个 run」，区间覆盖到的后续 run
    只删除被覆盖的字符。这样每个 run 各自的格式（加粗/颜色/字体）都能保住。
    """
    end = off + len(seg)
    res, cursor = [], 0
    for s, e, new in chosen:
        if e <= off or s >= end:
            continue  # 与该 run 无交集
        ls = max(s, off) - off  # 交集在本 run 内的起止
        le = min(e, end) - off
        if ls > cursor:
            res.append(seg[cursor:ls])  # 区间之前的原文
        if s >= off:  # 区间起点在本 run -> 新文本写这里
            res.append(new)
        cursor = max(cursor, le)  # 非起点 run：被覆盖的字符直接丢弃
    res.append(seg[cursor:])
    return "".join(res)


def replace_in_paragraph(p, pairs):
    """在一个段落里做跨 run 安全的多次替换，返回命中次数。"""
    nodes = _t_nodes(p)
    if not nodes:
        return 0
    segs = [n.text or "" for n in nodes]
    full = "".join(segs)
    if not full:
        return 0

    spans = []
    for old, new in pairs:
        if not old:
            continue
        start = 0
        while True:
            i = full.find(old, start)
            if i < 0:
                break
            spans.append((i, i + len(old), new))
            start = i + len(old)
    if not spans:
        return 0

    # 去重叠：起点升序、长度降序，贪心取不重叠的区间
    spans.sort(key=lambda t: (t[0], -(t[1] - t[0])))
    chosen, last_end = [], -1
    for sp in spans:
        if sp[0] < last_end:
            continue
        chosen.append(sp)
        last_end = sp[1]

    off = 0
    n = 0
    for node, seg in zip(nodes, segs):
        new_seg = _node_new_text(seg, off, chosen)
        if new_seg != seg:
            _set_t(node, new_seg)
            n += 1
        off += len(seg)
    return len(chosen)


def iter_paragraphs(doc, scope):
    """按范围产出 (位置标签, 段落对象)。scope 含 body / tables / header / footer。"""
    if "body" in scope:
        for i, p in enumerate(doc.paragraphs):
            yield "正文[%d]" % i, p
    if "tables" in scope:
        for ti, t in enumerate(doc.tables):
            for ri, row in enumerate(t.rows):
                for ci, cell in enumerate(row.cells):
                    for pi, p in enumerate(cell.paragraphs):
                        yield "表%d[%d,%d]段%d" % (ti, ri, ci, pi), p
    if "header" in scope:
        for si, s in enumerate(doc.sections):
            for pi, p in enumerate(s.header.paragraphs):
                yield "页眉[节%d,段%d]" % (si, pi), p
    if "footer" in scope:
        for si, s in enumerate(doc.sections):
            for pi, p in enumerate(s.footer.paragraphs):
                yield "页脚[节%d,段%d]" % (si, pi), p


# ---------------------------------------------------------------- 各项操作


def _first_paragraph_text(p, limit=40):
    t = p.text.replace("\n", " ")
    return t[:limit] + ("…" if len(t) > limit else "")


def cmd_list(doc):
    print("sections: %d" % len(doc.sections))
    for i, s in enumerate(doc.sections):
        print(
            "  [%d] 页眉=%r 页脚=%r"
            % (
                i,
                s.header.paragraphs[0].text if s.header.paragraphs else "",
                s.footer.paragraphs[0].text if s.footer.paragraphs else "",
            )
        )
    print("tables: %d" % len(doc.tables))
    for i, t in enumerate(doc.tables):
        head = " | ".join(c.text for c in t.rows[0].cells) if t.rows else ""
        print("  [%d] %dx%d  首行: %s" % (i, len(t.rows), len(t.columns), head[:50]))
    print("paragraphs: %d" % len(doc.paragraphs))
    for i, p in enumerate(doc.paragraphs):
        mark = " [图]" if _has_graphic(p) else ""
        print("  %3d| %s%s" % (i, _first_paragraph_text(p), mark))


def cmd_replace(doc, pairs, scope):
    hits, skipped = 0, 0
    for tag, p in iter_paragraphs(doc, scope):
        if _has_graphic(p):
            skipped += 1
            continue
        n = replace_in_paragraph(p, pairs)
        if n:
            hits += n
            print("   %s  %d 处" % (tag, n))
    print("replace: %d 处" % hits)
    if skipped:
        print("[note] 跳过 %d 个含图段落（契约要求不破坏图）" % skipped)
    if not hits:
        print("[warn] 一条都没命中：检查文字是否跨 run 或被切分，先用 --list 看实际文本")
    return hits


def _anchors(doc, needle, all_mode):
    """定位锚点段落。

    命中多处时默认拒绝执行：目录被 Word 刷成静态文本后，"1.2 资质与业绩" 这类
    标题在目录和正文里各有一份，静默地全部改一遍会把内容插进目录。
    宁可让人多敲一个 --all-anchors，也不能改错地方。
    """
    hits = [(i, p) for i, p in enumerate(doc.paragraphs) if needle in p.text]
    if not hits:
        sys.exit("找不到锚点段落：%r（用 --list 看实际文本）" % needle[:40])
    if len(hits) > 1 and not all_mode:
        print("[warn] 锚点 %r 命中 %d 段：" % (needle[:30], len(hits)))
        for i, p in hits:
            print("   %3d| %s" % (i, _first_paragraph_text(p)))
        # 用 print + exit(1) 而不是 sys.exit("...")：后者走 stderr，
        # 会和上面 stdout 的候选列表交错，输出顺序看着像乱了。
        print("请换更精确的锚点，或加 --all-anchors 对全部命中生效")
        sys.exit(1)
    return hits


def _clone_paragraph(anchor, text, style_name, doc, before):
    """复制锚点段落的 XML 作为新段落，继承其样式与 run 格式，再换成指定文字。"""
    new_p = copy.deepcopy(anchor._p)
    for tag in ("w:hyperlink", "w:bookmarkStart", "w:bookmarkEnd"):
        for el in new_p.findall(qn(tag)):
            new_p.remove(el)
    runs = new_p.findall(qn("w:r"))
    for extra in runs[1:]:
        new_p.remove(extra)
    if runs:
        run = runs[0]
        # 清掉 run 里除属性外的全部内容（图片/域/分隔符），只留 rPr 以保住格式
        for child in list(run):
            if child.tag != qn("w:rPr"):
                run.remove(child)
    else:
        run = OxmlElement("w:r")
        new_p.append(run)
    t = OxmlElement("w:t")
    run.append(t)
    _set_t(t, text)
    if before:
        anchor._p.addprevious(new_p)
    else:
        anchor._p.addnext(new_p)

    from docx.text.paragraph import Paragraph

    new_para = Paragraph(new_p, anchor._parent)
    if style_name:
        try:
            new_para.style = style_name
        except KeyError:
            names = [s.name for s in doc.styles if s.name]
            sys.exit("样式不存在：%s\n可用样式：%s" % (style_name, "、".join(names[:40])))
    return new_para


def cmd_insert(doc, needle, text, style_name, before, all_mode):
    hits = _anchors(doc, needle, all_mode)
    for i, p in hits:
        lines = text.split("\\n")
        # addnext 是插到锚点之后，逆序插入才能保证多行顺序正确；
        # addprevious 顺序插入天然正确。
        for line in lines if before else reversed(lines):
            _clone_paragraph(p, line, style_name, doc, before)
        print("   [%d] %r -> 插入 %d 段" % (i, needle[:20], len(lines)))
    print("insert: 命中 %d 个锚点" % len(hits))


def cmd_delete(doc, needle, all_mode):
    hits = _anchors(doc, needle, all_mode)
    for i, p in hits:
        print("   [%d] 删除 %r" % (i, p.text[:40]))
        p._p.getparent().remove(p._p)
    print("delete: %d 段" % len(hits))


def _cell(doc, ti, ri, ci):
    try:
        return doc.tables[ti].rows[ri].cells[ci]
    except IndexError:
        sys.exit("单元格越界：表%d[%d,%d]（共 %d 个表）" % (ti, ri, ci, len(doc.tables)))


def _find_template_rpr(table):
    """找表里第一个带格式 run 的 rPr（深拷贝），给空白格写入时借格式用。

    直接 p.add_run(value) 会掉回样式默认字体，填空白表单时整张表字体不统一。
    """
    for row in table.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    rpr = r._r.find(qn("w:rPr"))
                    if rpr is not None:
                        return copy.deepcopy(rpr)
    return None


def _set_cell_text(cell, value, table=None):
    """改单元格文字：保留首段首 run 的格式，其余段落删除。

    首段无 run（空白格）时，从 table 里借一个现成的 rPr，避免字体回退。
    """
    paras = cell.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    p = paras[0]
    runs = p.runs
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)
    if runs:
        for child in list(runs[0]._r):
            if child.tag != qn("w:rPr"):
                runs[0]._r.remove(child)
        t = OxmlElement("w:t")
        runs[0]._r.append(t)
        _set_t(t, value)
    else:
        r = p.add_run("")
        t = OxmlElement("w:t")
        r._r.append(t)
        _set_t(t, value)
        if table is not None and r._r.find(qn("w:rPr")) is None:
            rpr = _find_template_rpr(table)
            if rpr is not None:
                r._r.insert(0, rpr)


def cmd_cell(doc, ti, ri, ci, value):
    cell = _cell(doc, ti, ri, ci)
    old = cell.text
    _set_cell_text(cell, value, doc.tables[ti])
    print("cell 表%d[%d,%d]: %r -> %r" % (ti, ri, ci, old[:30], value))


def cmd_add_row(doc, ti, values):
    try:
        t = doc.tables[ti]
    except IndexError:
        sys.exit("表序号越界：%d（共 %d 个表）" % (ti, len(doc.tables)))
    row = t.add_row()
    for i, v in enumerate(values):
        if i < len(row.cells):
            _set_cell_text(row.cells[i], v, t)
    print("add-row 表%d: 新增 1 行（%d 列）" % (ti, len(row.cells)))


def _clear_row_text(tr):
    """清掉一行的可见文字（保留 rPr / tcPr / 域代码），用作空白模板行。"""
    for t in tr.findall(".//" + qn("w:t")):
        t.text = ""


def cmd_add_rows(doc, ti, n, template_row=None):
    """批量加行：deepcopy 模板行的 <w:tr>，边框/底纹/字号/对齐全部继承。

    python-docx 的 add_row() 产出的是裸行（丢格式），批量加行只能复制 XML。
    默认插在模板行之后；不指定模板行时复制最后一行。
    """
    try:
        t = doc.tables[ti]
    except IndexError:
        sys.exit("表序号越界：%d（共 %d 个表）" % (ti, len(doc.tables)))
    if not t.rows:
        sys.exit("表 %d 是空表，没有格式可复制" % ti)
    src_idx = template_row if template_row is not None else len(t.rows) - 1
    if not 0 <= src_idx < len(t.rows):
        sys.exit("模板行越界：%d（表 %d 共 %d 行）" % (src_idx, ti, len(t.rows)))
    src_tr = t.rows[src_idx]._tr
    anchor = src_tr
    for _ in range(n):
        new_tr = copy.deepcopy(src_tr)
        _clear_row_text(new_tr)
        anchor.addnext(new_tr)
        anchor = new_tr
    print("add-rows 表%d: 复制行%d格式 × %d 行" % (ti, src_idx, n))


def cmd_fill(doc, spec) -> int:
    """批量填表：values 逐格写入 表 ti 的 [start_row+i][start_col+j]。

    null 跳过（不清空现值）；合并单元格区按「写主格、跳过后续坐标」处理
    （python-docx 对 vMerge/gridSpan 返回同一个 cell 对象，靠 tc 身份去重）。
    """
    specs = spec if isinstance(spec, list) else [spec]
    total = 0
    for sp in specs:
        ti = sp.get("table", 0)
        r0 = sp.get("start_row", 0)
        c0 = sp.get("start_col", 0)
        if ti >= len(doc.tables):
            print("[warn] fill: 表 %d 越界，跳过" % ti)
            continue
        t = doc.tables[ti]
        seen = []  # 持有 tc 元素引用，保证 lxml 代理稳定、身份比较可靠
        for i, rowvals in enumerate(sp.get("values", [])):
            ri = r0 + i
            if ri >= len(t.rows):
                print("[warn] fill: 表 %d 行 %d 越界，该行跳过" % (ti, ri))
                continue
            cells = t.rows[ri].cells
            for j, v in enumerate(rowvals):
                if v is None:
                    continue
                ci = c0 + j
                if ci >= len(cells):
                    print("[warn] fill: 表 %d 行 %d 列 %d 越界，跳过" % (ti, ri, ci))
                    continue
                cell = cells[ci]
                if any(cell._tc is tc for tc in seen):
                    continue  # 合并区：主格已写，其余坐标不重复写
                seen.append(cell._tc)
                _set_cell_text(cell, str(v), t)
                total += 1
    print("fill: 写入 %d 格" % total)
    return total


def cmd_del_row(doc, ti, ri):
    try:
        t = doc.tables[ti]
        row = t.rows[ri]
    except IndexError:
        sys.exit("行列越界：表%d 行%d" % (ti, ri))
    row._tr.getparent().remove(row._tr)
    print("del-row 表%d 行%d: 已删除" % (ti, ri))


def _pick_sections(doc, sel):
    """选节：body=最后一节（正文节，避开封面），all=全部，数字=指定序号。"""
    n = len(doc.sections)
    if sel == "all":
        return list(range(n))
    if sel == "body":
        return [n - 1]
    try:
        i = int(sel)
    except ValueError:
        sys.exit("--section 只能是 all / body / 数字，收到 %r" % sel)
    if not 0 <= i < n:
        sys.exit("节序号越界：%d（共 %d 节）" % (i, n))
    return [i]


def _set_hf_text(container, text):
    """改页眉/页脚文字：先断开与上一节的链接，否则改不到或会连带改别的节。"""
    if container.is_linked_to_previous:
        container.is_linked_to_previous = False
    paras = container.paragraphs
    if not paras:
        paras = [container.add_paragraph()]
    p = paras[0]
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    runs = p.runs
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)
    if runs:
        for child in list(runs[0]._r):
            if child.tag != qn("w:rPr"):
                runs[0]._r.remove(child)
        t = OxmlElement("w:t")
        runs[0]._r.append(t)
        _set_t(t, text)
    else:
        p.add_run(text)


def _load_fill(path):
    """--fill 数据文件：.json 按 schema 解析；.csv 读成 values 矩阵（空格跳过）。"""
    import csv
    import io
    import json

    with open(path, encoding="utf-8-sig") as f:
        text = f.read()
    if path.lower().endswith(".csv"):
        rows = [r for r in csv.reader(io.StringIO(text)) if r]
        return {"table": 0, "values": [[v if v != "" else None for v in r] for r in rows]}
    return json.loads(text)


def cmd_hf(doc, text, sel, is_header):
    what = "页眉" if is_header else "页脚"
    for i in _pick_sections(doc, sel):
        s = doc.sections[i]
        c = s.header if is_header else s.footer
        _set_hf_text(c, text)
        print("%s 节%d: -> %r" % (what, i, text))


def cmd_verify(path):
    """复用生成链路的验收：让 Word 真正打开一次。打不开 = 结构改坏了。"""
    tmp = tempfile.mkdtemp(prefix="texere_edit_")
    try:
        pdf = os.path.join(tmp, "verify.pdf")
        r = subprocess.run(
            [sys.executable, os.path.join(KIT, "scripts", "finalize.py"), path, pdf],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            # 子进程管道统一 UTF-8，避免 GBK→utf-8 解码乱码（同 render.py）
            env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        )
        print(r.stdout.strip())
        if r.returncode != 0:
            print(r.stderr.strip()[:1500])
            sys.exit("验收失败：Word 打不开改后的文档，结构可能已损坏（备份仍在）")
        print("verify: OK")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------- 入口


def main():
    ap = argparse.ArgumentParser(add_help=True)
    ap.add_argument("docx")
    ap.add_argument("--list", action="store_true", help="只打印结构，不改文件")
    ap.add_argument("--replace", action="append", default=[], metavar="旧=新")
    ap.add_argument(
        "--scope",
        default="body,tables",
        help="--replace 的作用范围，逗号分隔：body,tables,header,footer（默认 body,tables）",
    )
    ap.add_argument("--after", metavar="锚点文字")
    ap.add_argument("--before", metavar="锚点文字")
    ap.add_argument("--text", help="配合 --after/--before；\\n 表示另起一段")
    ap.add_argument("--style", help="新段落样式名（不存在则报错，不静默降级）")
    ap.add_argument("--delete", metavar="段落所含文字")
    ap.add_argument(
        "--all-anchors",
        action="store_true",
        help="锚点/删除命中多处时仍然全部执行（默认拒绝并列出候选）",
    )
    ap.add_argument("--cell", nargs=4, metavar=("表", "行", "列", "值"))
    ap.add_argument("--add-row", nargs="+", metavar=("表", "值"))
    ap.add_argument(
        "--add-rows",
        nargs=2,
        metavar=("表", "数量"),
        help="批量加空行：复制模板行的全部格式（边框/底纹/字号），默认复制最后一行",
    )
    ap.add_argument(
        "--template-row", type=int, default=None, metavar="行", help="配合 --add-rows：复制哪一行"
    )
    ap.add_argument(
        "--fill",
        action="append",
        default=[],
        metavar="data.json",
        help='批量填表：[{"table":0,"start_row":1,"values":[[...]]}] 或 CSV，可多次',
    )
    ap.add_argument("--del-row", nargs=2, metavar=("表", "行"))
    ap.add_argument("--header")
    ap.add_argument("--footer")
    ap.add_argument(
        "--section",
        default="body",
        help="页眉页脚作用于哪些节：body(默认,最后一节) / all / 序号",
    )
    ap.add_argument("--out", help="另存为（默认写回原文件）")
    ap.add_argument("--no-backup", action="store_true")
    ap.add_argument("--verify", action="store_true", help="改完调 Word 打开一次做验收")
    a = ap.parse_args()

    if not os.path.exists(a.docx):
        sys.exit("找不到文件: " + a.docx)

    doc = Document(a.docx)

    if a.list:
        cmd_list(doc)
        return

    changed = False

    if a.replace:
        pairs = []
        for item in a.replace:
            if "=" not in item:
                sys.exit("--replace 要写成 旧=新，收到 %r" % item)
            old, new = item.split("=", 1)
            pairs.append((old, new))
        scope = [s.strip() for s in a.scope.split(",") if s.strip()]
        changed |= cmd_replace(doc, pairs, scope) > 0

    if a.after or a.before:
        if not a.text:
            sys.exit("--after/--before 必须配 --text")
        cmd_insert(doc, a.after or a.before, a.text, a.style, bool(a.before), a.all_anchors)
        changed = True

    if a.delete:
        cmd_delete(doc, a.delete, a.all_anchors)
        changed = True

    if a.cell:
        cmd_cell(doc, int(a.cell[0]), int(a.cell[1]), int(a.cell[2]), a.cell[3])
        changed = True

    if a.add_row:
        cmd_add_row(doc, int(a.add_row[0]), a.add_row[1:])
        changed = True

    if a.add_rows:
        cmd_add_rows(doc, int(a.add_rows[0]), int(a.add_rows[1]), a.template_row)
        changed = True

    for fill_path in a.fill:
        if not os.path.exists(fill_path):
            sys.exit("找不到 fill 数据文件: " + fill_path)
        cmd_fill(doc, _load_fill(fill_path))
        changed = True

    if a.del_row:
        cmd_del_row(doc, int(a.del_row[0]), int(a.del_row[1]))
        changed = True

    if a.header:
        cmd_hf(doc, a.header, a.section, True)
        changed = True

    if a.footer:
        cmd_hf(doc, a.footer, a.section, False)
        changed = True

    if not changed:
        print("没有任何改动（未写盘）。用 --list 看结构，或检查参数。")
        return

    out = a.out or a.docx
    if not a.out and not a.no_backup:
        bak = os.path.splitext(a.docx)[0] + ".bak.docx"
        shutil.copy2(a.docx, bak)
        print("backup ->", bak)
    doc.save(out)
    print("saved:", out)

    if a.verify:
        cmd_verify(out)


if __name__ == "__main__":
    main()
