"""编辑链共享的 docx 修改原语：edit.py 与 patch.py 共用，纯函数、无 CLI 副作用。

历史：这些原语曾寄居在 edit.py 的私有函数里（_anchors/_clone_paragraph/_set_cell_text），
patch.py 靠 sys.path hack + importlib 伸手拿人，还要用 `except (Exception, SystemExit)`
把 edit 的 sys.exit 兜住——「退出进程」成了跨模块协议的一部分，set_cell 越界时
SystemExit 更是直接穿透 except 把整次 patch 打崩。现在原语只抛 DocxEditError
子类（str(e) 即 CLI 可见措辞，与历史文案逐字一致）；退出码、歧义候选列表等
由 edit.py 的 CLI 层决定，patch 路径的失败由 except Exception 收成结构化报告。
"""

import copy

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


class DocxEditError(Exception):
    """编辑原语层的错误基类；消息就是 CLI 可见措辞，调用方不要改写。"""


class AnchorNotFound(DocxEditError):  # noqa: N818  # 名字是计划/文档定下的公开契约
    pass


class StyleNotFound(DocxEditError):  # noqa: N818
    pass


class CellOutOfRange(DocxEditError):  # noqa: N818
    pass


# ---------------------------------------------------------------- 基础工具


def t_nodes(p):
    """段落内全部 w:t 节点（含超链接里的文字），按文档顺序。"""
    return p._p.findall(".//" + qn("w:t"))


def set_t(node, text):
    """写 w:t 文本；首尾有空白时必须加 xml:space="preserve"，否则 Word 会吞掉空格。"""
    node.text = text
    if text[:1].isspace() or text[-1:].isspace():
        node.set(XML_SPACE, "preserve")
    else:
        node.attrib.pop(XML_SPACE, None)


def has_graphic(p):
    """段落里含图片 / 图形对象时返回 True。契约要求不破坏图，遇到就跳过。"""
    for tag in ("w:drawing", "w:pict", "w:object"):
        if p._p.find(".//" + qn(tag)) is not None:
            return True
    return False


def node_new_text(seg, off, chosen):
    """按原有 run 切分，算出这个 run 替换后应该是什么文本。

    seg     : 该 run 原本的文本
    off     : 该 run 在整段文本中的起始偏移
    chosen  : 已去重叠的替换区间 [(start, end, new), ...]

    关键点：替换内容只写进「区间起点所在的那个 run」，区间覆盖到的后续 run
    只删除被覆盖的字符。这样每个 run 各自的格式（加粗/颜色/字体）都能保住。
    """
    end = off + len(seg)
    res, cursor = [], 0
    for s, e, new in chosen:
        if e <= off or s >= end:
            continue  # 与该 run 无交集
        ls = max(s, off) - off  # 交集在本 run 内的起止
        le = min(e, end) - off
        if ls > cursor:
            res.append(seg[cursor:ls])  # 区间之前的原文
        if s >= off:  # 区间起点在本 run -> 新文本写这里
            res.append(new)
        cursor = max(cursor, le)  # 非起点 run：被覆盖的字符直接丢弃
    res.append(seg[cursor:])
    return "".join(res)


def replace_in_paragraph(p, pairs):
    """在一个段落里做跨 run 安全的多次替换，返回命中次数。"""
    nodes = t_nodes(p)
    if not nodes:
        return 0
    segs = [n.text or "" for n in nodes]
    full = "".join(segs)
    if not full:
        return 0

    spans = []
    for old, new in pairs:
        if not old:
            continue
        start = 0
        while True:
            i = full.find(old, start)
            if i < 0:
                break
            spans.append((i, i + len(old), new))
            start = i + len(old)
    if not spans:
        return 0

    # 去重叠：起点升序、长度降序，贪心取不重叠的区间
    spans.sort(key=lambda t: (t[0], -(t[1] - t[0])))
    chosen, last_end = [], -1
    for sp in spans:
        if sp[0] < last_end:
            continue
        chosen.append(sp)
        last_end = sp[1]

    off = 0
    n = 0
    for node, seg in zip(nodes, segs):
        new_seg = node_new_text(seg, off, chosen)
        if new_seg != seg:
            set_t(node, new_seg)
            n += 1
        off += len(seg)
    return len(chosen)


def iter_paragraphs(doc, scope):
    """按范围产出 (位置标签, 段落对象)。scope 含 body / tables / header / footer。"""
    if "body" in scope:
        for i, p in enumerate(doc.paragraphs):
            yield "正文[%d]" % i, p
    if "tables" in scope:
        for ti, t in enumerate(doc.tables):
            for ri, row in enumerate(t.rows):
                for ci, cell in enumerate(row.cells):
                    for pi, p in enumerate(cell.paragraphs):
                        yield "表%d[%d,%d]段%d" % (ti, ri, ci, pi), p
    if "header" in scope:
        for si, s in enumerate(doc.sections):
            for pi, p in enumerate(s.header.paragraphs):
                yield "页眉[节%d,段%d]" % (si, pi), p
    if "footer" in scope:
        for si, s in enumerate(doc.sections):
            for pi, p in enumerate(s.footer.paragraphs):
                yield "页脚[节%d,段%d]" % (si, pi), p


# ---------------------------------------------------------------- 锚点与克隆


def find_anchors(doc, needle):
    """返回文本含 needle 的正文段落 [(索引, 段落), ...]；零命中抛 AnchorNotFound。

    歧义（多处命中）的处置是 CLI 决策——是否拒绝、是否列候选，留给调用方；
    patch 的 insert op 历来取全部命中，直接吃返回值即可。
    """
    hits = [(i, p) for i, p in enumerate(doc.paragraphs) if needle in p.text]
    if not hits:
        raise AnchorNotFound("找不到锚点段落：%r（用 --list 看实际文本）" % needle[:40])
    return hits


def clone_paragraph(anchor, text, style_name, doc, before):
    """复制锚点段落的 XML 作为新段落，继承其样式与 run 格式，再换成指定文字。

    样式不存在抛 StyleNotFound（附可用样式清单）——不静默降级。
    """
    new_p = copy.deepcopy(anchor._p)
    for tag in ("w:hyperlink", "w:bookmarkStart", "w:bookmarkEnd"):
        for el in new_p.findall(qn(tag)):
            new_p.remove(el)
    runs = new_p.findall(qn("w:r"))
    for extra in runs[1:]:
        new_p.remove(extra)
    if runs:
        run = runs[0]
        # 清掉 run 里除属性外的全部内容（图片/域/分隔符），只留 rPr 以保住格式
        for child in list(run):
            if child.tag != qn("w:rPr"):
                run.remove(child)
    else:
        run = OxmlElement("w:r")
        new_p.append(run)
    t = OxmlElement("w:t")
    run.append(t)
    set_t(t, text)
    if before:
        anchor._p.addprevious(new_p)
    else:
        anchor._p.addnext(new_p)

    new_para = Paragraph(new_p, anchor._parent)
    if style_name:
        try:
            new_para.style = style_name
        except KeyError:
            names = [s.name for s in doc.styles if s.name]
            raise StyleNotFound(
                "样式不存在：%s\n可用样式：%s" % (style_name, "、".join(names[:40]))
            )
    return new_para


# ---------------------------------------------------------------- 单元格与行


def cell_at(doc, ti, ri, ci):
    """取 表ti[ri,ci] 单元格；越界抛 CellOutOfRange（消息含表总数）。"""
    try:
        return doc.tables[ti].rows[ri].cells[ci]
    except IndexError:
        raise CellOutOfRange(
            "单元格越界：表%d[%d,%d]（共 %d 个表）" % (ti, ri, ci, len(doc.tables))
        )


def find_template_rpr(table):
    """找表里第一个带格式 run 的 rPr（深拷贝），给空白格写入时借格式用。

    直接 p.add_run(value) 会掉回样式默认字体，填空白表单时整张表字体不统一。
    """
    for row in table.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    rpr = r._r.find(qn("w:rPr"))
                    if rpr is not None:
                        return copy.deepcopy(rpr)
    return None


def set_cell_text(cell, value, table=None):
    """改单元格文字：保留首段首 run 的格式，其余段落删除。

    首段无 run（空白格）时，从 table 里借一个现成的 rPr，避免字体回退。
    """
    paras = cell.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    p = paras[0]
    runs = p.runs
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)
    if runs:
        for child in list(runs[0]._r):
            if child.tag != qn("w:rPr"):
                runs[0]._r.remove(child)
        t = OxmlElement("w:t")
        runs[0]._r.append(t)
        set_t(t, value)
    else:
        r = p.add_run("")
        t = OxmlElement("w:t")
        r._r.append(t)
        set_t(t, value)
        if table is not None and r._r.find(qn("w:rPr")) is None:
            rpr = find_template_rpr(table)
            if rpr is not None:
                r._r.insert(0, rpr)


def clear_row_text(tr):
    """清掉一行的可见文字（保留 rPr / tcPr / 域代码），用作空白模板行。"""
    for t in tr.findall(".//" + qn("w:t")):
        t.text = ""
