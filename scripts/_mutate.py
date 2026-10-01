"""mutate 能力的纯逻辑核（可导入、无 CLI 副作用）。

从 patch.py 外提声明式 Patch 的执行内核：schema 校验、前置条件检查、7 个操作原语
（统一签名 `op(doc, spec) -> MutationResult`）、以及 assess / dry_run / apply /
validate 编排。底层写操作继续走 `_docx_edit`（docx kernel）。patch.py 作为 thin
CLI 壳 re-export 同名符号，并把 `MutationResult` 降级回历史一致的 `tuple[bool, str]`
与 stdout/stderr 措辞；argparse / sys.exit / 证据写盘仍留壳。

裸名 sibling import，禁止反向 import patch / edit。
"""

from dataclasses import dataclass, field

import _docx_edit
from _shared import sha256_file
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


@dataclass
class MutationResult:
    """单个操作的结构化结论。

    op(doc, spec) 统一返回它；壳层用 `.ok` / `.message` 降级回 `tuple[bool, str]`。
    `evidence` 留给需要携带额外取证信息（坐标、前后值）的操作，当前编排链只消费
    ok / message，保证与原实现逐字节等价。
    """

    ok: bool
    message: str = ""
    evidence: dict = field(default_factory=dict)


# =============================================================================
# Patch Schema Validation
# =============================================================================


def validate_patch_schema(patch: dict) -> tuple[bool, list[str]]:
    """验证 Patch JSON 是否符合 schema。"""
    errors = []

    # 必需字段
    if "id" not in patch:
        errors.append("缺少必需字段：id")
    if "operations" not in patch:
        errors.append("缺少必需字段：operations")
    elif not isinstance(patch["operations"], list):
        errors.append("operations 必须是数组")

    # 验证每个 operation
    for i, op in enumerate(patch.get("operations", [])):
        if "op" not in op:
            errors.append(f"operation[{i}] 缺少必需字段：op")
            continue

        valid_ops = {
            "replace_text",
            "insert_after",
            "insert_before",
            "delete_paragraph",
            "set_cell",
            "add_row",
            "del_row",
        }
        if op["op"] not in valid_ops:
            errors.append(f"operation[{i}] 无效的 op: {op['op']}")

        if "target" not in op:
            errors.append(f"operation[{i}] 缺少必需字段：target")

    # 验证 preconditions（可选）
    if "preconditions" in patch:
        pc = patch["preconditions"]
        if "hash" in pc and not isinstance(pc["hash"], str):
            errors.append("preconditions.hash 必须是字符串")
        if "must_contain" in pc and not isinstance(pc["must_contain"], list):
            errors.append("preconditions.must_contain 必须是数组")

    if errors:
        import sys

        print("Patch schema validation failed:", file=sys.stderr)
        for error in errors:
            print(f"  - {error}", file=sys.stderr)
        return False, errors
    return True, []


# =============================================================================
# Hash Precondition
# =============================================================================


def compute_docx_hash(docx_path: str) -> str:
    """计算 DOCX 文件的 SHA256 hash（实现收敛到 _shared.sha256_file）。"""
    digest = sha256_file(docx_path)
    if digest is None:
        raise FileNotFoundError(f"无法计算 hash，文件不存在：{docx_path}")
    return digest


def check_hash_precondition(docx_path: str, expected_hash: str) -> tuple[bool, str]:
    """检查文档 hash 是否匹配。"""
    actual = compute_docx_hash(docx_path)
    if actual == expected_hash:
        return True, "Hash 校验通过"
    else:
        return (
            False,
            f"Hash 不匹配 (期望:{expected_hash[:16]}..., 实际:{actual[:16]}...)",
        )


def check_must_contain_precondition(doc: Document, must_contain: list[str]) -> tuple[bool, str]:
    """检查文档是否包含所有必需文本。"""
    missing = []
    for text in must_contain:
        found = False
        for para in doc.paragraphs:
            if text in para.text:
                found = True
                break
        if not found:
            missing.append(text)

    if missing:
        return False, f"缺少必需文本：{missing}"
    return True, "所有必需文本都存在"


# =============================================================================
# Operations Implementation（统一 op(doc, spec) -> MutationResult）
# =============================================================================


def _edit_module():
    """共享编辑原语模块（历史上这里做过 sys.path hack + importlib 拿 edit.py
    私有函数；原语已下沉到 _docx_edit.py，正常兄弟导入即可）。保留函数只为
    测试/调用方有一个稳定的取模块入口。"""
    return _docx_edit


def replace_text_op(doc: Document, op: dict) -> MutationResult:
    """替换文本操作。"""
    target = op["target"]
    expected_old = op.get("expected_old_text")
    new_text = op["new_text"]

    para_idx = target.get("paragraph")
    if para_idx is None:
        return MutationResult(False, "target.paragraph 是必需的")

    try:
        para = doc.paragraphs[para_idx]
        old_text = para.text

        if expected_old and expected_old not in old_text:
            return MutationResult(
                False, f"预期旧文本不存在：{expected_old} (实际：{old_text[:50]})"
            )

        # 使用共享编辑原语的跨 run 替换逻辑
        # 构造 pairs
        pairs = [(expected_old or old_text, new_text)]
        n = _docx_edit.replace_in_paragraph(para, pairs)

        if n > 0:
            return MutationResult(True, f"成功替换 (第{para_idx}段)")
        else:
            return MutationResult(False, "未找到匹配文本")
    except IndexError:
        return MutationResult(False, f"段落索引越界：{para_idx}")
    except Exception as e:
        return MutationResult(False, f"替换失败：{e}")


def insert_after_op(doc: Document, op: dict) -> MutationResult:
    """在锚点后插入操作。"""
    target = op["target"]
    content = op["content"]
    style = op.get("style", "Normal")

    anchor_text = target.get("anchor")
    if not anchor_text:
        return MutationResult(False, "target.anchor 是必需的")

    try:
        # patch 的 insert 取全部命中（写补丁时已人工确认过锚点）；
        # 零命中抛 AnchorNotFound，由下面 except 收成结构化「插入失败」
        hits = _docx_edit.find_anchors(doc, anchor_text)

        for para_idx, anchor_para in hits:
            # addnext 每次插在锚点紧邻之后，逆序遍历使最终文档顺序与 content 一致
            for line in reversed(content):
                _docx_edit.clone_paragraph(anchor_para, line, style, doc, before=False)

        return MutationResult(True, f"在 {len(hits)} 个位置插入 {len(content)} 段")
    except Exception as e:
        return MutationResult(False, f"插入失败：{e}")


def insert_before_op(doc: Document, op: dict) -> MutationResult:
    """在锚点前插入操作。"""
    target = op["target"]
    content = op["content"]
    style = op.get("style", "Normal")

    anchor_text = target.get("anchor")
    if not anchor_text:
        return MutationResult(False, "target.anchor 是必需的")

    try:
        hits = _docx_edit.find_anchors(doc, anchor_text)

        for para_idx, anchor_para in hits:
            # addprevious 天然保持正序，顺序遍历即可
            for line in content:
                _docx_edit.clone_paragraph(anchor_para, line, style, doc, before=True)

        return MutationResult(True, f"在 {len(hits)} 个位置插入 {len(content)} 段")
    except Exception as e:
        return MutationResult(False, f"插入失败：{e}")


def delete_paragraph_op(doc: Document, op: dict) -> MutationResult:
    """删除段落操作。"""
    target = op["target"]

    para_idx = target.get("paragraph")
    if para_idx is None:
        return MutationResult(False, "target.paragraph 是必需的")

    try:
        para = doc.paragraphs[para_idx]
        para._p.getparent().remove(para._p)
        return MutationResult(True, f"已删除第{para_idx}段")
    except IndexError:
        return MutationResult(False, f"段落索引越界：{para_idx}")
    except Exception as e:
        return MutationResult(False, f"删除失败：{e}")


def set_cell_op(doc: Document, op: dict) -> MutationResult:
    """设置单元格值操作。"""
    target = op["target"]
    value = op["value"]

    try:
        ti = target.get("table")
        ri = target.get("row")
        ci = target.get("column")

        if None in (ti, ri, ci):
            return MutationResult(False, "target 需要 table, row, column")

        cell = _docx_edit.cell_at(doc, ti, ri, ci)
        old = cell.text
        _docx_edit.set_cell_text(cell, value, doc.tables[ti])

        return MutationResult(True, f"单元格 [{ti}][{ri},{ci}]: {old[:20]} -> {value[:20]}")
    except Exception as e:
        return MutationResult(False, f"设置单元格失败：{e}")


def _insert_row(table, after_row):
    """在 after_row（0 基）之后插入一行并返回它；after_row 为 None/负数时追加到表尾。

    python-docx 的 add_row() 只会追加，且产出的是裸行（丢边框/底纹/字号）。
    要让 after_row 真正生效，走 edit.py cmd_add_rows 的路子：deepcopy 锚点行的
    <w:tr> 继承全部格式，清掉文字，再 addnext 插到锚点之后。
    """
    import copy

    from docx.table import _Row

    rows = table.rows
    n = len(rows)
    if after_row is None or after_row < 0:
        return table.add_row()
    if after_row >= n:
        raise IndexError("after_row %d 越界（表共 %d 行）" % (after_row, n))

    src_tr = rows[after_row]._tr
    new_tr = copy.deepcopy(src_tr)
    _docx_edit.clear_row_text(new_tr)
    src_tr.addnext(new_tr)
    return _Row(new_tr, table)


def add_row_op(doc: Document, op: dict) -> MutationResult:
    """添加行操作：插在 target.after_row（0 基）之后，缺省追加到表尾。

    新行克隆锚点行的格式，文字先清空再按 values 写入；未给值的单元格留空，
    不会残留锚点行的旧文字。
    """
    target = op["target"]
    values = op.get("values", [])

    try:
        ti = target.get("table")
        if ti is None:
            return MutationResult(False, "target.table 是必需的")
        if not 0 <= ti < len(doc.tables):
            return MutationResult(False, f"表格索引越界：{ti}（共 {len(doc.tables)} 个表）")

        table = doc.tables[ti]
        after_row = target.get("after_row")
        row = _insert_row(table, after_row)

        for i, v in enumerate(values):
            if i < len(row.cells):
                _docx_edit.set_cell_text(row.cells[i], v, table)

        where = "表尾" if (after_row is None or after_row < 0) else f"第{after_row}行之后"
        return MutationResult(True, f"表{ti}: 在{where}新增 1 行（{len(values)}列）")
    except Exception as e:
        return MutationResult(False, f"添加行失败：{e}")


def del_row_op(doc: Document, op: dict) -> MutationResult:
    """删除行操作。"""
    target = op["target"]

    try:
        ti = target.get("table")
        ri = target.get("row")

        if None in (ti, ri):
            return MutationResult(False, "target 需要 table 和 row")

        # 手动实现删除
        table = doc.tables[ti]
        row = table.rows[ri]
        row._tr.getparent().remove(row._tr)

        return MutationResult(True, f"表{ti}行{ri}: 已删除")
    except Exception as e:
        return MutationResult(False, f"删除行失败：{e}")


# Operation registry
OPERATIONS = {
    "replace_text": replace_text_op,
    "insert_after": insert_after_op,
    "insert_before": insert_before_op,
    "delete_paragraph": delete_paragraph_op,
    "set_cell": set_cell_op,
    "add_row": add_row_op,
    "del_row": del_row_op,
}


# =============================================================================
# Assessment & Dry Run
# =============================================================================


def assess_patch(patch: dict, doc: Document) -> tuple[bool, list[str]]:
    """评估 Patch 可行性，不进行实际修改。"""
    issues = []

    for i, op in enumerate(patch.get("operations", [])):
        op_name = op.get("op")
        target = op.get("target", {})

        # 静态检查
        if op_name == "replace_text":
            para_idx = target.get("paragraph")
            if para_idx is not None and para_idx >= len(doc.paragraphs):
                issues.append(f"operation[{i}]: 段落索引越界")

        elif op_name in ("insert_after", "insert_before"):
            anchor = target.get("anchor")
            if anchor:
                count = sum(1 for p in doc.paragraphs if anchor in p.text)
                if count == 0:
                    issues.append(f"operation[{i}]: 锚点不存在：{anchor}")

        elif op_name == "set_cell":
            ti, ri = target.get("table"), target.get("row")  # ci 预留参数
            if ti is not None and ti >= len(doc.tables):
                issues.append(f"operation[{i}]: 表格索引越界")
            elif ti is not None:
                if ri is not None and ri >= len(doc.tables[ti].rows):
                    issues.append(f"operation[{i}]: 行索引越界")

        elif op_name == "add_row":
            ti, ar = target.get("table"), target.get("after_row")
            if ti is not None and ti < len(doc.tables):
                n_rows = len(doc.tables[ti].rows)
                if ar is not None and ar >= n_rows:
                    issues.append(f"operation[{i}]: after_row 越界：{ar}（表共 {n_rows} 行）")

    return len(issues) == 0, issues


def dry_run_patch(patch: dict, doc: Document) -> tuple[bool, str]:
    """Dry run：模拟执行 Patch，预估影响。"""
    import copy

    # 深拷贝文档进行模拟
    doc_copy = copy.deepcopy(doc)

    failed_ops = []
    successful_ops = []

    for i, op in enumerate(patch.get("operations", [])):
        op_name = op.get("op")

        try:
            if op_name in OPERATIONS:
                res = OPERATIONS[op_name](doc_copy, op)
                if res.ok:
                    successful_ops.append((i, op_name, res.message))
                else:
                    failed_ops.append((i, op_name, res.message))
            else:
                failed_ops.append((i, op_name, "未实现的 operation"))
        except Exception as e:
            failed_ops.append((i, op_name, "异常：%s" % e))

    if failed_ops:
        # 只回报结果，不在这里 sys.exit：本函数签名是 (bool, str)，调用方拿到
        # False 才能自己决定退出码；内部抢先 exit 会让调用方的 else 分支成死代码。
        detail = "\n".join("  [%d] %s: %s" % (i, op, msg) for i, op, msg in failed_ops)
        return False, "Dry run 失败：%d 个操作不可执行\n%s" % (len(failed_ops), detail)
    return (
        True,
        "Dry run 成功：%d 个操作将通过\n" % len(successful_ops)
        + "\n".join("  [%d] %s: %s" % (i, op, msg) for i, op, msg in successful_ops),
    )


# =============================================================================
# Apply & Validate
# =============================================================================


def apply_patch(patch: dict, doc: Document) -> tuple[bool, list[str]]:
    """应用 Patch 到文档。"""
    errors = []

    for i, op in enumerate(patch.get("operations", [])):
        op_name = op.get("op")

        if op_name not in OPERATIONS:
            errors.append(f"operation[{i}]: 未实现的 operation: {op_name}")
            continue

        res = OPERATIONS[op_name](doc, op)
        if not res.ok:
            errors.append(f"operation[{i}] ({op_name}): {res.message}")

    return len(errors) == 0, errors


def validate_patch_result(doc: Document, patch: dict) -> tuple[bool, list[str]]:
    """验证 Patch 应用后的结果。"""
    issues = []

    # 验证每个 operation 是否达到预期
    for i, op in enumerate(patch.get("operations", [])):
        op_name = op.get("op")

        if op_name == "replace_text":
            expected = op.get("new_text")
            target = op["target"]
            para_idx = target.get("paragraph")

            if para_idx < len(doc.paragraphs):
                if expected not in doc.paragraphs[para_idx].text:
                    issues.append(f"operation[{i}]: 预期文本未找到：{expected}")

        elif op_name == "set_cell":
            target = op["target"]
            expected = op["value"]
            ti, ri, ci = target.get("table"), target.get("row"), target.get("column")

            if ti is not None and ri is not None and ci is not None:
                if expected not in doc.tables[ti].rows[ri].cells[ci].text:
                    issues.append(f"operation[{i}]: 单元格值未更新")

        elif op_name == "add_row":
            # after_row 是新加的能力，验证也要跟上：新行必须落在声明的位置
            target = op.get("target", {})
            ti, ar = target.get("table"), target.get("after_row")
            values = op.get("values") or []
            if ti is not None and ti < len(doc.tables) and values:
                rows = doc.tables[ti].rows
                idx = (ar + 1) if (ar is not None and ar >= 0) else len(rows) - 1
                if 0 <= idx < len(rows):
                    row_text = "".join(c.text for c in rows[idx].cells)
                    if values[0] not in row_text:
                        issues.append(f"operation[{i}]: 第{idx}行未写入预期值：{values[0]}")

    return len(issues) == 0, issues


# =============================================================================
# Edit-side helpers（从 edit.py 收拢的非 CLI 纯逻辑；cmd_* 的 print/sys.exit 留壳）
# =============================================================================


def first_paragraph_text(p, limit=40):
    """截断段落文字做单行预览（edit.py --list 与歧义锚点列表共用）。"""
    t = p.text.replace("\n", " ")
    return t[:limit] + ("…" if len(t) > limit else "")


def set_header_footer_text(container, text):
    """改页眉/页脚文字：先断开与上一节的链接，否则改不到或会连带改别的节。"""
    if container.is_linked_to_previous:
        container.is_linked_to_previous = False
    paras = container.paragraphs
    if not paras:
        paras = [container.add_paragraph()]
    p = paras[0]
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    runs = p.runs
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)
    if runs:
        for child in list(runs[0]._r):
            if child.tag != qn("w:rPr"):
                runs[0]._r.remove(child)
        t = OxmlElement("w:t")
        runs[0]._r.append(t)
        _docx_edit.set_t(t, text)
    else:
        p.add_run(text)
